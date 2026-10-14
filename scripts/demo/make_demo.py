#!/usr/bin/env python3
"""Makes the made half of Biblo's demo files, in Resources/Demo.noindex. Run it with scripts/make-demo.sh.

Everything in them is generic: public facts about Microsoft Foundry, Fabric and Copilot, the fictitious Contoso, and
made-up numbers. Pictures are HTML drawn by a headless Edge or Chrome, voices are macOS's `say`, videos ffmpeg, decks
python-pptx, Word and RTF files textutil. The text and code files next to them are written by hand and left alone,
but every file gets its date from DATES (a demo where everything is from today looks like one).

Logos (scripts/demo/assets, see its SOURCES.md) appear only as their owners allow: whole, unaltered and unanimated, in
documents, diagrams and slides; never recolored, distorted, given effects or used as Biblo's own."""
from __future__ import annotations

import argparse
import datetime as dt
import html
import io
import itertools
import math
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pygments import lex
from pygments.lexers import PythonLexer
from pygments.styles import get_style_by_name

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[1] / "Resources" / "Demo.noindex"
ASSETS = HERE / "assets"
MADE = {".pdf", ".docx", ".rtf", ".pptx", ".png", ".jpg", ".m4a", ".mp4", ".mov"}
FFMPEG = shutil.which("ffmpeg") or "/opt/homebrew/bin/ffmpeg"
WORK = Path(tempfile.mkdtemp(prefix="biblo-demo-"))
BROWSER = ""
BUILDERS: dict[str, callable] = {}
_pages = itertools.count()


def builder(function):
    BUILDERS[function.__name__] = function
    return function


def run(*args, **options) -> str:
    result = subprocess.run([str(a) for a in args], capture_output=True, text=True, **options)
    if result.returncode != 0:
        sys.exit(f"failed: {' '.join(map(str, args))}\n{result.stderr[-3000:]}")
    return result.stdout


def find_browser() -> str:
    for path in [os.environ.get("BROWSER", ""), "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
                 "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                 "/Applications/Chromium.app/Contents/MacOS/Chromium"]:
        if path and os.access(path, os.X_OK):
            return path
    sys.exit("needs Microsoft Edge, Google Chrome or Chromium (or BROWSER=/path/to/one)")


def out(relative: str) -> Path:
    path = OUT / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def asset(name: str) -> str:
    return (ASSETS / name).as_uri()


def esc(text: str) -> str:
    return html.escape(text, quote=False)


# MARK: Drawing


def page(body: str, css: str = "", size: tuple[int, int] = (1280, 720), bg: str = "#fff") -> str:
    w, h = size
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
*{{box-sizing:border-box}}
html,body{{margin:0;width:{w}px;height:{h}px;overflow:hidden;background:{bg}}}
body{{font-family:system-ui,"Helvetica Neue",Helvetica,Arial,sans-serif;color:#0f172a;
-webkit-font-smoothing:antialiased;text-rendering:geometricPrecision}}
{css}</style></head><body>{body}</body></html>"""


def render(document: str, size: tuple[int, int], scale: float = 1, transparent: bool = False) -> Image.Image:
    """A page drawn by the browser: `size` CSS pixels, times `scale` in the picture."""
    source = WORK / f"page{next(_pages)}.html"
    source.write_text(document, encoding="utf-8")
    shot = WORK / "shot.png"
    shot.unlink(missing_ok=True)
    arguments = [BROWSER, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--no-first-run",
                 "--no-default-browser-check", "--run-all-compositor-stages-before-draw",
                 f"--user-data-dir={WORK / 'browser'}", f"--window-size={size[0]},{size[1]}",
                 f"--force-device-scale-factor={scale}", f"--screenshot={shot}"]
    if transparent:
        arguments.append("--default-background-color=00000000")
    run(*arguments, source.as_uri())
    image = Image.open(shot)
    image.load()
    return image


def print_pdf(document: str, path: Path) -> None:
    source = WORK / f"page{next(_pages)}.html"
    source.write_text(document, encoding="utf-8")
    run(BROWSER, "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
        f"--user-data-dir={WORK / 'browser'}", "--no-pdf-header-footer", f"--print-to-pdf={path}", source.as_uri())


def save_png(image: Image.Image, path: Path) -> None:
    image.save(path, optimize=True)


def save_jpg(image: Image.Image, path: Path, quality: int = 86) -> None:
    image.convert("RGB").save(path, quality=quality, optimize=True, progressive=True)


def textutil(document: str, path: Path, title: str) -> None:
    """A Word or RTF file from HTML. The author is set: textutil would put the Mac's owner there."""
    source = WORK / f"page{next(_pages)}.html"
    source.write_text(document, encoding="utf-8")
    run("textutil", "-convert", path.suffix[1:], "-title", title, "-author", "Contoso", "-output", path, source)


_logos: dict[tuple, Path] = {}


def logo(svg: str, side: int = 1024, fill: float = 0.8, bg: str = "transparent") -> Path:
    """An official logo as a PNG, scaled evenly and otherwise as it is (scripts/demo/assets/SOURCES.md)."""
    key = (svg, side, fill, bg)
    if key not in _logos:
        width = int(side * fill)
        document = page(f'<img src="{asset(svg)}" style="width:{width}px;height:auto">',
                        "body{display:grid;place-items:center}", (side, side), bg)
        path = WORK / f"logo{len(_logos)}.png"
        render(document, (side, side), transparent=bg == "transparent").save(path)
        _logos[key] = path
    return _logos[key]


def window(content: str, title: str, size: tuple[int, int], dark: bool, css: str = "",
           toolbar: str = "") -> str:
    """A Mac window with its shadow, as ⇧⌘4 then Space takes it."""
    w, h = size
    margin = 56
    chrome = "#2b2b2d" if dark else "#ececee"
    border = "rgba(255,255,255,.14)" if dark else "rgba(0,0,0,.18)"
    title_color = "#d6d6d8" if dark else "#3c3c43"
    return page(f"""<div class="win"><div class="bar"><i style="background:#ff5f57"></i><i style="background:#febc2e">
</i><i style="background:#28c840"></i><span>{esc(title)}</span>{toolbar}</div><div class="content">{content}</div></div>""",
                f""".win{{position:absolute;left:{margin}px;top:{margin - 16}px;width:{w}px;height:{h}px;border-radius:12px;
overflow:hidden;box-shadow:0 22px 70px rgba(0,0,0,.45),0 0 0 .5px {border};display:flex;flex-direction:column}}
.bar{{height:38px;background:{chrome};display:flex;align-items:center;gap:8px;padding:0 14px;position:relative;
flex:none;border-bottom:1px solid {'#000' if dark else '#d4d4d8'}}}
.bar i{{width:12px;height:12px;border-radius:50%;box-shadow:inset 0 0 0 .5px rgba(0,0,0,.2)}}
.bar span{{position:absolute;left:0;right:0;text-align:center;font:600 13px system-ui;color:{title_color};
pointer-events:none}}
.content{{flex:1;position:relative;overflow:hidden}}
{css}""", (w + 2 * margin, h + 2 * margin), "transparent")


# MARK: Sound and video


def say(text: str, voice: str, rate: int | None = None) -> Path:
    path = WORK / f"voice{next(_pages)}.aiff"
    run("say", "-v", voice, *(["-r", rate] if rate else []), "-o", path, text)
    return path


def seconds(path: Path) -> float:
    return float(run(FFMPEG.replace("ffmpeg", "ffprobe"), "-v", "error", "-show_entries", "format=duration",
                     "-of", "csv=p=0", path))


def join_audio(parts: list[tuple[Path, float]], path: Path, bitrate: str = "48k") -> None:
    """Voices one after the other, each followed by its pause (seconds), as a mono AAC file."""
    inputs, chains = [], []
    for index, (part, pause) in enumerate(parts):
        inputs += ["-i", part]
        chains.append(f"[{index}:a]aresample=24000,aformat=channel_layouts=mono,apad=pad_dur={pause}[a{index}]")
    labels = "".join(f"[a{i}]" for i in range(len(parts)))
    run(FFMPEG, "-v", "error", "-y", *inputs, "-filter_complex",
        ";".join(chains) + f";{labels}concat=n={len(parts)}:v=0:a=1,loudnorm=I=-18:TP=-2[out]", "-map", "[out]",
        "-map_metadata", "-1", "-c:a", "aac", "-b:a", bitrate, "-movflags", "+faststart", path)


def slideshow(slides: list[tuple[Image.Image, Path | None]], path: Path, fade: float = 0.5,
              hold: float = 2.5, pad: float = 0.7) -> None:
    """Narrated slides: each stays on while its voice (if any) speaks, then dissolves into the next."""
    frames, durations = [], []
    for index, (image, voice) in enumerate(slides):
        frame = WORK / f"slide{next(_pages)}.png"
        image.convert("RGB").save(frame)
        frames.append(frame)
        durations.append((seconds(voice) + pad if voice else hold) + fade)
    inputs, offsets, time = [], [], 0.0
    for frame, duration in zip(frames, durations):
        inputs += ["-loop", "1", "-framerate", "25", "-t", f"{duration:.3f}", "-i", frame]
    for duration in durations[:-1]:
        time += duration - fade
        offsets.append(time)
    chain, last = [], "[0:v]"
    for index, offset in enumerate(offsets, start=1):
        chain.append(f"{last}[{index}:v]xfade=transition=fade:duration={fade}:offset={offset:.3f}[v{index}]")
        last = f"[v{index}]"
    chain.append(f"{last}format=yuv420p[video]")
    # Each voice starts with its slide: the slides before it, less the dissolves.
    voices = [(voice, index) for index, (_, voice) in enumerate(slides) if voice]
    audio_inputs, audio_chain = [], []
    for number, (voice, index) in enumerate(voices):
        start = sum(durations[:index]) - fade * index + (0.25 if index else 0.4)
        audio_inputs += ["-i", voice]
        audio_chain.append(f"[{len(frames) + number}:a]aresample=24000,aformat=channel_layouts=mono,"
                           f"adelay={int(start * 1000)}[s{number}]")
    if voices:
        mixed = "".join(f"[s{n}]" for n in range(len(voices)))
        audio_chain.append(f"{mixed}amix=inputs={len(voices)}:normalize=0,loudnorm=I=-18:TP=-2[audio]")
    total = sum(durations) - fade * (len(durations) - 1)
    run(FFMPEG, "-v", "error", "-y", *inputs, *audio_inputs, "-filter_complex", ";".join(chain + audio_chain),
        "-map", "[video]", *(["-map", "[audio]", "-c:a", "aac", "-b:a", "64k"] if voices else []),
        "-t", f"{total:.3f}", "-c:v", "libx264", "-preset", "slow", "-tune", "stillimage", "-crf", "30", "-r", "25",
        "-map_metadata", "-1", "-movflags", "+faststart", path)


def recording(states: list[tuple[Image.Image, float]], path: Path) -> None:
    """A screen recording: still states, cut from one to the next."""
    listing = WORK / f"states{next(_pages)}.txt"
    lines = []
    for image, duration in states:
        frame = WORK / f"state{next(_pages)}.png"
        image.convert("RGB").save(frame)
        lines += [f"file '{frame}'", f"duration {duration}"]
    lines.append(f"file '{frame}'")
    listing.write_text("\n".join(lines) + "\n")
    run(FFMPEG, "-v", "error", "-y", "-f", "concat", "-safe", "0", "-i", listing, "-vf",
        "fps=30,format=yuv420p", "-c:v", "libx264", "-preset", "slow", "-tune", "stillimage", "-crf", "26",
        "-map_metadata", "-1", "-movflags", "+faststart", path)


def perspective(image: Image.Image, corners: list[tuple[float, float]]) -> Image.Image:
    """The picture as if photographed: its corners moved to `corners` (top left, top right, bottom right, bottom
    left, as fractions of its size)."""
    w, h = image.size
    target = [(x * w, y * h) for x, y in corners]
    source = [(0, 0), (w, 0), (w, h), (0, h)]
    # PIL maps each output pixel back to the input: solve for the 8 coefficients that take target to source.
    rows, values = [], []
    for (x, y), (u, v) in zip(target, source):
        rows.append([x, y, 1, 0, 0, 0, -u * x, -u * y]); values.append(u)
        rows.append([0, 0, 0, x, y, 1, -v * x, -v * y]); values.append(v)
    n = 8
    matrix = [row + [value] for row, value in zip(rows, values)]
    for column in range(n):
        pivot = max(range(column, n), key=lambda r: abs(matrix[r][column]))
        matrix[column], matrix[pivot] = matrix[pivot], matrix[column]
        for r in range(n):
            if r != column:
                factor = matrix[r][column] / matrix[column][column]
                matrix[r] = [a - factor * b for a, b in zip(matrix[r], matrix[column])]
    coefficients = [matrix[r][n] / matrix[r][r] for r in range(n)]
    return image.transform(image.size, Image.Transform.PERSPECTIVE, coefficients, Image.Resampling.BICUBIC,
                           fillcolor=(160, 150, 134))


def photographed(image: Image.Image, seed: int) -> Image.Image:
    """Light falloff, a little blur and sensor noise: a phone's photo of something flat."""
    rng = random.Random(seed)
    w, h = image.size
    image = image.convert("RGB").filter(ImageFilter.GaussianBlur(0.9))
    shade = Image.new("L", (w // 8, h // 8), 0)
    draw = ImageDraw.Draw(shade)
    cx, cy = w // 16 + rng.randint(-20, 20), h // 18
    for step in range(60, 0, -1):
        radius = step * max(w, h) / 8 / 60 * 1.25
        draw.ellipse([cx - radius, cy - radius, cx + radius, cy + radius], fill=int(255 * (1 - step / 60) ** 0.6))
    shade = shade.resize((w, h), Image.Resampling.BICUBIC).filter(ImageFilter.GaussianBlur(40))
    dark = Image.new("RGB", (w, h), (70, 62, 52))
    image = Image.composite(image, Image.blend(image, dark, 0.38), shade)
    noise = Image.effect_noise((w, h), 9).convert("RGB")
    image = Image.blend(image, noise, 0.035)
    warm = Image.new("RGB", (w, h), (255, 236, 205))
    return Image.blend(image, warm, 0.06)


@builder
def pictures():
    """Images/: logos, a chart and a whiteboard photo."""
    shutil.copy(logo("microsoft-fabric.svg"), out("Images/Fabric logo.png"))
    shutil.copy(logo("microsoft-foundry.svg"), out("Images/Foundry logo.png"))
    shutil.copy(logo("microsoft-copilot-2023.svg"), out("Images/Copilot logo (2023 icon).png"))
    shutil.copy(logo("github-copilot.svg", fill=0.62, bg="#fff"), out("Images/GitHub Copilot icon.png"))
    lockup = page(f'<img src="{asset("github-copilot-lockup.svg")}" style="width:1240px;height:auto">',
                  "body{display:grid;place-items:center}", (1600, 520), "#fff")
    save_png(render(lockup, (1600, 520)), out("Images/GitHub Copilot logo.png"))

    teams = [("Customer Support", 58, 81), ("Sales", 61, 78), ("Engineering", 52, 72), ("Marketing", 47, 66),
             ("Finance", 33, 49), ("People & HR", 29, 44), ("Legal", 14, 27)]
    rows = "".join(f"""<div class="row"><div class="team">{team}</div><div class="track">
<div class="bar before" style="width:{q2}%"></div><div class="bar now" style="width:{q3}%"><b>{q3}%</b></div></div>
<div class="delta">+{q3 - q2}</div></div>""" for team, q2, q3 in teams)
    chart = page(f"""<main><p class="eyebrow">Contoso Ltd. &nbsp;·&nbsp; Workplace AI pulse &nbsp;·&nbsp; Sample data</p>
<h1>Copilot adoption by team</h1><p class="sub">Weekly active users as a share of licensed seats, Q2 → Q3 2026</p>
<div class="legend"><span><i class="before"></i>Q2</span><span><i class="now"></i>Q3</span><span class="pts">pts</span></div>
<section>{rows}</section><footer><span>Fictitious company and numbers, made for a product demo.</span>
<span>n = 1,240 seats</span></footer></main>""", """
main{position:absolute;inset:64px 84px}
.eyebrow{margin:0;font:600 17px/1 system-ui;letter-spacing:.14em;text-transform:uppercase;color:#64748b}
h1{margin:22px 0 10px;font:800 64px/1.05 system-ui;letter-spacing:-.025em;color:#0b1020}
.sub{margin:0;font:400 25px/1.3 system-ui;color:#475569}
.legend{display:flex;gap:30px;margin:44px 0 18px 250px;font:600 19px system-ui;color:#334155}
.legend span{display:flex;align-items:center;gap:10px}.legend i{width:30px;height:14px;border-radius:7px}
.legend .pts{margin-left:auto;margin-right:6px;color:#94a3b8;font-weight:500}
i.before,.bar.before{background:#e2e8f0}
i.now,.bar.now{background:linear-gradient(90deg,#2563eb,#7c3aed 55%,#db2777)}
.row{display:grid;grid-template-columns:250px 1fr 70px;align-items:center;height:76px;border-top:1px solid #eef2f7}
.team{font:600 23px system-ui;color:#1e293b}
.track{position:relative;height:44px}
.bar{position:absolute;left:0;border-radius:22px}.bar.before{top:4px;height:36px}
.bar.now{top:10px;height:24px;display:flex;justify-content:flex-end;align-items:center;padding-right:12px}
.bar.now b{font:700 16px system-ui;color:#fff}
.delta{text-align:right;font:700 21px system-ui;color:#059669}
footer{position:absolute;bottom:0;left:0;right:0;display:flex;justify-content:space-between;
font:400 17px system-ui;color:#94a3b8;border-top:1px solid #e2e8f0;padding-top:18px}""", (1600, 1000))
    save_png(render(chart, (1600, 1000)), out("Images/Copilot adoption by team – Q3 2026.png"))

    board = page("""<div class="wall"><div class="board">
<h1>AGENT IDEAS — Q4</h1>
<ol class="ideas">
<li>Returns helper <span class="arrow">→</span> <u>Foundry agent</u> + order lookup tool</li>
<li>Weekly sales recap from <u>OneLake</u> data</li>
<li>Onboarding buddy (HR FAQ)</li>
<li>Meeting notes → action items</li>
</ol>
<svg class="flow" viewBox="0 0 900 420">
<g fill="none" stroke="#1f2937" stroke-width="6" stroke-linecap="round" stroke-linejoin="round">
<path d="M40 70 C 38 40, 210 36, 214 64 C 218 98, 44 104, 40 70 Z"/>
<path d="M330 40 L 560 34 L 566 120 L 334 126 Z"/>
<path d="M690 70 C 688 34, 880 34, 884 68 C 888 104, 694 108, 690 70 Z"/>
<path d="M224 80 L 318 82 M 300 66 L 320 82 L 298 96"/>
<path d="M572 80 L 676 78 M 656 62 L 678 78 L 656 94"/>
<path d="M450 132 C 452 190, 448 230, 452 268 M 436 250 L 452 272 L 470 250"/>
</g>
<g font-family="Marker Felt" font-size="40" fill="#1f2937" text-anchor="middle">
<text x="127" y="82">user</text><text x="450" y="92">AGENT</text><text x="787" y="82">tools</text>
</g>
<text x="452" y="330" font-family="Marker Felt" font-size="44" fill="#dc2626" text-anchor="middle">evals before launch!!</text>
<path d="M190 300 C 230 262, 700 258, 720 304 C 738 350, 220 372, 190 326 Z" fill="none" stroke="#dc2626"
stroke-width="5" stroke-linecap="round"/>
<text x="452" y="404" font-family="Marker Felt" font-size="36" fill="#15803d" text-anchor="middle">ground answers in our docs ✓</text>
</svg>
<div class="next"><b>NEXT</b><br>☐ book Foundry workshop (Oct 14)<br>☐ Fabric workspace access for the team<br>
☐ draft eval set — 50 real questions</div>
<div class="note yellow">Pilot:<br>50 users<br>4 weeks</div>
<div class="note pink">Ask Legal about data retention</div>
<div class="note blue">Measure: CSAT + time to resolve</div>
<div class="tray"><i class="m1"></i><i class="m2"></i><i class="m3"></i></div>
</div></div>""", """
.wall{position:absolute;inset:0;background:radial-gradient(circle at 40% 30%,#d9d2c5,#a99f8f)}
.board{position:absolute;left:110px;top:90px;right:110px;bottom:150px;border-radius:10px;
background:radial-gradient(ellipse at 35% 25%,#ffffff 0%,#f6f7f5 45%,#e6e8e4 100%);
border:22px solid #c9ccd0;box-shadow:inset 0 0 0 3px #9ca3af,0 30px 60px rgba(0,0,0,.35)}
.board:before{content:"";position:absolute;inset:0;background:
radial-gradient(ellipse at 70% 60%,rgba(120,130,140,.10),transparent 40%),
radial-gradient(ellipse at 20% 80%,rgba(120,130,140,.08),transparent 35%)}
h1{position:absolute;left:120px;top:70px;margin:0;font:400 118px "Marker Felt";color:#1d4ed8;transform:rotate(-1.6deg)}
.ideas{position:absolute;left:110px;top:270px;width:1000px;margin:0;padding-left:60px;font:400 54px/1.5 "Marker Felt";
color:#111827;transform:rotate(-.8deg)}
.ideas li{padding-left:10px}.ideas u{text-decoration-thickness:4px;text-underline-offset:10px;text-decoration-color:#2563eb}
.arrow{color:#2563eb}
.next{position:absolute;left:150px;top:880px;font:400 44px/1.45 "Marker Felt";color:#15803d;transform:rotate(-1.2deg)}
.next b{font-weight:400;color:#1f2937;text-decoration:underline;text-decoration-thickness:4px}
.flow{position:absolute;right:120px;top:110px;width:900px;height:420px;transform:rotate(1.2deg)}
.note{position:absolute;width:250px;height:230px;padding:26px 24px;font:400 34px/1.25 Noteworthy;color:#1f2937;
box-shadow:0 10px 18px rgba(0,0,0,.22)}
.yellow{background:#fde68a;right:640px;bottom:120px;transform:rotate(-4deg)}
.pink{background:#fbcfe8;right:350px;bottom:100px;transform:rotate(3deg)}
.blue{background:#bfdbfe;right:70px;bottom:130px;transform:rotate(-2deg)}
.tray{position:absolute;left:-22px;right:-22px;bottom:-60px;height:38px;background:linear-gradient(#d1d5db,#9ca3af);
border-radius:0 0 8px 8px}
.tray i{position:absolute;top:-14px;width:150px;height:22px;border-radius:11px}
.m1{left:300px;background:#1f2937}.m2{left:480px;background:#dc2626}.m3{left:660px;background:#2563eb}""",
                 (2400, 1600))
    photo = perspective(render(board, (2400, 1600)),
                        [(-0.02, 0.035), (1.02, -0.025), (1.025, 1.02), (-0.015, 0.985)])
    photo = photographed(photo, 7).resize((2000, 1333), Image.Resampling.LANCZOS)
    save_jpg(photo, out("Images/IMG_4127.jpg"), 84)


GLYPHS = {
    "person": '<circle cx="24" cy="16" r="8"/><path d="M8 42c2-10 9-15 16-15s14 5 16 15"/>',
    "chat": '<path d="M8 10h32v22H22l-9 8v-8H8z"/><path d="M15 19h18M15 25h11"/>',
    "search": '<circle cx="21" cy="21" r="11"/><path d="M29 29l11 11"/>',
    "pulse": '<path d="M6 26h9l5-12 8 22 5-10h9"/>',
    "check": '<rect x="8" y="8" width="32" height="32" rx="7"/><path d="M16 24l6 6 11-12"/>',
}


def glyph(name: str, color: str) -> str:
    return (f'<svg class="glyph" viewBox="0 0 48 48" fill="none" stroke="{color}" stroke-width="3" '
            f'stroke-linecap="round" stroke-linejoin="round">{GLYPHS[name]}</svg>')


@builder
def diagram():
    """Images/: an architecture diagram that uses the Foundry and Fabric icons, as their terms allow."""
    def card(x, y, w, h, body, accent, extra=""):
        return (f'<div class="card {extra}" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px;'
                f'--accent:{accent}">{body}</div>')

    def head(icon, name, detail):
        return f'<div class="head">{icon}<div><h3>{name}</h3><p>{detail}</p></div></div>'

    def badge(x, y, number):
        return (f'<g><circle cx="{x}" cy="{y}" r="17" fill="#0f172a"/><text x="{x}" y="{y + 6}" fill="#fff" '
                f'font-size="17" font-weight="700" text-anchor="middle">{number}</text></g>')

    foundry = f'<img class="logo" src="{asset("microsoft-foundry.svg")}" style="width:84px;height:84px">'
    fabric = f'<img class="logo" src="{asset("microsoft-fabric.svg")}" style="width:64px;height:auto">'
    rows = [("Agent Service", "runs the <b>returns-assistant</b> agent"),
            ("Models", "gpt-5-mini · Global Standard"),
            ("Tools", "<code>lookup_order()</code> · <code>create_return()</code> · file&nbsp;search"),
            ("Knowledge", "policy PDFs in a File Search vector store"),
            ("Identity", "managed identity, no keys in code"),
            ("Guardrails", "on prompts, outputs and tool calls"),
            ("Tracing", "every run, step and tool call")]
    foundry_body = (head(foundry, "Microsoft Foundry", "Project <b>proj-support</b> · Sweden Central")
                    + '<dl>' + "".join(f"<dt>{a}</dt><dd>{b}</dd>" for a, b in rows) + "</dl>")
    cards = "".join([
        card(72, 470, 250, 150, head(glyph("person", "#475569"), "Customers", "web and mobile"), "#94a3b8"),
        card(400, 455, 310, 180, head(glyph("chat", "#0891b2"), "Web chat",
                                      "React app on Azure Container Apps"), "#0891b2"),
        card(790, 262, 580, 560, foundry_body, "#4f46e5", "big"),
        card(1470, 262, 380, 170, head(glyph("search", "#2563eb"), "File Search",
                                       "Vector store: return policy, shipping options"), "#2563eb"),
        card(1470, 457, 380, 170, head(fabric, "Microsoft Fabric", "OneLake lakehouse: orders, returns"),
             "#0f9d8a"),
        card(1470, 652, 380, 170, head(glyph("pulse", "#d97706"), "Application Insights",
                                       "Traces, latency and tokens per run"), "#d97706"),
        card(790, 878, 580, 112, head(glyph("check", "#16a34a"), "Evaluations",
                                      "Groundedness · relevance · tool-call accuracy · safety"), "#16a34a"),
    ])
    steps = ["Customer asks a question in the web chat.", "File search retrieves passages from the policy PDFs.",
             "It calls <code>lookup_order()</code> on OneLake data.",
             "It answers with citations; the run is traced.", "Evaluations gate every release."]
    legend = "".join(f"<li><b>{n}</b><span>{t}</span></li>" for n, t in enumerate(steps, start=1))
    wires = f"""<svg class="wires" viewBox="0 0 1920 1080">
<defs><marker id="tip" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="9" markerHeight="9" orient="auto">
<path d="M0 0L10 5L0 10z" fill="#334155"/></marker></defs>
<g stroke="#334155" stroke-width="2.5" fill="none" marker-end="url(#tip)">
<path d="M322 545H396"/><path d="M710 520H784"/><path d="M790 575H716"/>
<path d="M1370 347H1464"/><path d="M1370 542H1464"/><path d="M1370 737H1464" stroke-dasharray="7 6"/>
<path d="M1080 822V872" stroke-dasharray="7 6"/></g>
{badge(747, 498, 1)}{badge(1417, 323, 2)}{badge(1417, 518, 3)}{badge(747, 600, 4)}{badge(1110, 848, 5)}</svg>"""
    document = page(f"""<header><p class="eyebrow">Architecture · draft v0.3 · Contoso Outdoor</p>
<h1>Returns assistant on Microsoft Foundry + Microsoft Fabric</h1>
<p class="sub">Retrieval-augmented agent: answers from our policies, order data from OneLake.</p></header>
{wires}{cards}<ol class="legend">{legend}</ol>
<footer>Sample diagram for a product demo. Contoso is a fictitious company. Icons: Azure Architecture Icons and
Microsoft Fabric icons, shown as provided.</footer>""", """
body{background:#fbfcfe;background-image:radial-gradient(#e2e8f0 1.2px,transparent 1.2px);background-size:28px 28px}
header{position:absolute;left:72px;top:56px}
.eyebrow{margin:0;font:600 16px system-ui;letter-spacing:.14em;text-transform:uppercase;color:#64748b}
h1{margin:16px 0 10px;font:800 46px/1.1 system-ui;letter-spacing:-.02em;color:#0b1020}
.sub{margin:0;font:400 22px system-ui;color:#475569}
.wires{position:absolute;inset:0;width:1920px;height:1080px;font-family:system-ui}
.card{position:absolute;background:#fff;border:1.5px solid #dbe3ee;border-radius:18px;padding:22px 24px;
box-shadow:0 1px 2px rgba(15,23,42,.05),0 8px 24px rgba(15,23,42,.06);border-top:5px solid var(--accent)}
.head{display:flex;gap:18px;align-items:center}
.head h3{margin:0;font:700 23px system-ui;color:#0f172a}
.head p{margin:6px 0 0;font:400 17px/1.35 system-ui;color:#475569}
.glyph{width:46px;height:46px;flex:none}
.big{padding:30px 34px}.big .head h3{font-size:32px}.big .head p{font-size:18px}
dl{display:grid;grid-template-columns:150px 1fr;margin:30px 0 0;border-top:1px solid #e2e8f0}
dt,dd{margin:0;padding:15px 0;border-bottom:1px solid #eef2f7;font:17px system-ui}
dt{font-weight:700;color:#3730a3}dd{color:#334155}
code{font:600 15.5px "SF Mono",Menlo,monospace;color:#0f172a;background:#f1f5f9;padding:2px 6px;border-radius:5px}
.legend{position:absolute;left:72px;top:690px;width:640px;margin:0;padding:22px 26px;list-style:none;
background:#fff;border:1.5px dashed #cbd5e1;border-radius:18px}
.legend li{display:flex;gap:14px;align-items:baseline;padding:7px 0;font:17px/1.35 system-ui;color:#334155}
.legend b{flex:none;display:grid;place-items:center;width:26px;height:26px;border-radius:50%;background:#0f172a;
color:#fff;font-size:14px}
footer{position:absolute;left:72px;right:72px;bottom:28px;font:14px system-ui;color:#94a3b8}""",
                    (1920, 1080))
    save_png(render(document, (1920, 1080)), out("Images/RAG architecture – Foundry + Fabric.png"))


# MARK: Screenshots

TERMINAL = """<span class="p">~/src/contoso-support-agent</span> <span class="m">main</span> <span class="g">❯</span> azd up

<span class="g">?</span> <b>Select an Azure Subscription to use:</b> <span class="c"> 1. Contoso Demo (00000000-0000-0000-0000-000000000000)</span>
<span class="g">?</span> <b>Enter a value for the 'location' infrastructure parameter:</b> <span class="c">14. (Europe) Sweden Central (swedencentral)</span>

<b>Packaging services (azd package)</b>

  <span class="g">(✓) Done:</span> Packaging service web
  <span class="d">- Image Hash: sha256:8c1f4b07e2d94e1d0a6f3b2c9e7a51d4c0b8f6e2a9d3c7b1e5f0a4d8c2b6e9f1</span>
  <span class="d">- Target Image: contoso-support-agent/web-dev:azd-deploy-1791470520</span>

<b>Provisioning Azure resources (azd provision)</b>
Provisioning Azure resources can take some time.

Subscription: <span class="c">Contoso Demo (00000000-0000-0000-0000-000000000000)</span>
Location: <span class="c">Sweden Central</span>

  You can view detailed progress in the Azure Portal:
  <span class="c u">https://portal.azure.com/#view/HubsExtension/DeploymentDetailsBlade/~/overview/id/%2Fsubscriptions%2F0000…</span>

  <span class="g">(✓) Done:</span> Resource group: <span class="c">rg-contoso-support-dev</span> <span class="d">(1.4s)</span>
  <span class="g">(✓) Done:</span> Log Analytics workspace: <span class="c">log-k3x7q</span> <span class="d">(22.6s)</span>
  <span class="g">(✓) Done:</span> Application Insights: <span class="c">appi-k3x7q</span> <span class="d">(5.1s)</span>
  <span class="g">(✓) Done:</span> Foundry: <span class="c">ai-contoso-k3x7q</span> <span class="d">(34.9s)</span>
  <span class="g">(✓) Done:</span> Foundry project: <span class="c">proj-support</span> <span class="d">(12.2s)</span>
  <span class="g">(✓) Done:</span> Container Registry: <span class="c">crk3x7q</span> <span class="d">(25.3s)</span>
  <span class="g">(✓) Done:</span> Container Apps Environment: <span class="c">cae-k3x7q</span> <span class="d">(1m33s)</span>
  <span class="g">(✓) Done:</span> Container App: <span class="c">ca-web-k3x7q</span> <span class="d">(19.8s)</span>

<b>Deploying services (azd deploy)</b>

  <span class="g">(✓) Done:</span> Deploying service web
  - Endpoint: <span class="c u">https://ca-web-k3x7q.wittyfield-3c9a1e52.swedencentral.azurecontainerapps.io/</span>

<span class="g">SUCCESS:</span> Your up workflow to provision and deploy to Azure completed in <b>6 minutes 12 seconds</b>.
<span class="p">~/src/contoso-support-agent</span> <span class="m">main</span> <span class="g">❯</span> <span class="cursor"> </span>"""

TERMINAL_CSS = """.content{background:#1c1c1e}
pre{margin:0;padding:14px 18px;font:13px/19px "SF Mono",Menlo,monospace;color:#e5e5e7;white-space:pre}
pre b{color:#fff}.g{color:#5fd068}.c{color:#5ac8e8}.m{color:#c792ea}.p{color:#7aa2f7;font-weight:700}
.d{color:#8e8e93}.u{text-decoration:underline;text-underline-offset:2px}
.cursor{background:#e5e5e7}"""

CHAT_CSS = """.bar{height:50px;background:#f3f3f5}
.nav{position:absolute;left:96px;top:15px;display:flex;gap:18px;color:#8e8e93;font:600 17px system-ui}
.url{position:absolute;left:50%;top:10px;transform:translateX(-50%);width:430px;height:30px;border-radius:8px;
background:#e4e4e8;display:flex;align-items:center;justify-content:center;font:500 13px system-ui;color:#1c1c1e}
.url i{display:inline-block;width:8px;height:8px;margin-right:8px;border-radius:50%;background:#34c759}
.app{position:absolute;inset:0;display:flex;flex-direction:column;background:#f6f5f1}
.top{flex:none;display:flex;align-items:center;gap:14px;height:68px;padding:0 36px;background:#fff;
border-bottom:1px solid #e7e5df}
.mark{width:34px;height:34px;border-radius:9px;background:#1f4d3a;position:relative;overflow:hidden}
.mark:before{content:"";position:absolute;left:5px;bottom:7px;border-left:12px solid transparent;
border-right:12px solid transparent;border-bottom:17px solid #f2c14e}
.mark:after{content:"";position:absolute;left:13px;bottom:7px;border-left:8px solid transparent;
border-right:8px solid transparent;border-bottom:12px solid #fff}
.brand{font:800 19px system-ui;letter-spacing:-.01em;color:#1f2a24}
.brand small{display:block;font:500 12.5px system-ui;color:#6b7280;letter-spacing:0}
.pill{margin-left:6px;padding:4px 10px;border-radius:999px;background:#e8f2ec;color:#1f4d3a;font:700 11.5px system-ui;
letter-spacing:.06em;text-transform:uppercase}
.note{margin-left:auto;font:12.5px system-ui;color:#6b7280}
.thread{flex:1;overflow:hidden;display:flex;flex-direction:column;justify-content:flex-end;gap:16px;
padding:24px 0 18px;width:760px;margin:0 auto}
.msg{max-width:600px;padding:14px 18px;border-radius:18px;font:15.5px/1.5 system-ui;color:#1f2a24}
.msg.user{align-self:flex-end;background:#1f4d3a;color:#fff;border-bottom-right-radius:6px}
.msg.agent{align-self:flex-start;background:#fff;border:1px solid #e7e5df;border-bottom-left-radius:6px;
box-shadow:0 1px 2px rgba(0,0,0,.04)}
.msg b{font-weight:700}
.who{display:flex;align-items:center;gap:8px;margin-bottom:6px;font:700 12.5px system-ui;color:#1f4d3a}
.who:before{content:"";width:8px;height:8px;border-radius:50%;background:#2f9e6a}
.cites{display:flex;gap:8px;margin-top:12px}
.cite{display:flex;align-items:center;gap:7px;padding:5px 10px;border-radius:8px;background:#f4f1ea;
font:600 12.5px system-ui;color:#5b5446}
.cite:before{content:"";width:10px;height:13px;border-radius:2px;border:1.5px solid #8a7f68;background:#fff}
.tool{align-self:flex-start;display:flex;align-items:center;gap:9px;padding:7px 12px;border-radius:10px;
background:#eef2ff;border:1px solid #dfe5ff;font:600 13px "SF Mono",Menlo,monospace;color:#3f3fb0}
.tool:before{content:"ƒ";font:italic 700 15px Georgia,serif;color:#5b5bd6}
.tool span{color:#6b6bb8;font-weight:500}
.tool.running:after{content:"• • •";letter-spacing:2px;color:#9a9ae0}
.chips{display:flex;gap:10px}
.chip{padding:8px 14px;border-radius:999px;border:1px solid #d9d4c7;background:#fff;font:600 13.5px system-ui;
color:#3d463f}
.composer{flex:none;width:760px;margin:0 auto 26px;display:flex;align-items:center;gap:12px;height:56px;
padding:0 10px 0 20px;border-radius:16px;background:#fff;border:1px solid #dcd8cd;
box-shadow:0 4px 16px rgba(31,42,36,.06);font:15.5px system-ui;color:#9ca3af}
.composer.typed{color:#1f2a24}
.composer .caret{display:inline-block;width:2px;height:20px;margin-left:1px;vertical-align:-4px;background:#1f4d3a}
.send{margin-left:auto;width:38px;height:38px;border-radius:11px;background:#d9d4c7;display:grid;place-items:center;
color:#fff;font:700 18px system-ui}
.composer.typed .send{background:#1f4d3a}"""

GREETING = ("agent", "Hi! I'm the Contoso Outdoor assistant. I can check your orders, start a return or help you "
                     "choose gear.", [])
ASK = "Hi! The Trailhead 2 tent I ordered arrived with a cracked pole. Can I return it?"
ANSWER = ("agent", "Sorry about that, Alex. I found order <b>#48213</b>: a Trailhead 2-person tent, delivered October "
                   "3. It's inside the 30-day return window, and damaged items ship back free.<br><br>I've opened "
                   "return <b>RMA-20931</b> and emailed you a prepaid label. Would you like a replacement sent now, "
                   "or a refund to your card?", ["Return policy.pdf · p. 2"])
ASK_AGAIN = "A replacement, please. I need it for a trip next weekend."
DONE = ("agent", "Done! A new Trailhead 2 ships today by express delivery and should arrive on Wednesday, October 14, "
                 "well before your trip. Keep the damaged tent until the label arrives.",
        ["Shipping options.pdf · p. 1"])


def chat(thread: list, draft: str = "", chips: bool = False) -> str:
    """The Contoso Outdoor support chat (a fictitious shop) in a browser, at one moment."""
    parts = []
    for item in thread:
        kind = item[0]
        if kind == "user":
            parts.append(f'<div class="msg user">{esc(item[1])}</div>')
        elif kind == "agent":
            cites = "".join(f'<span class="cite">{esc(c)}</span>' for c in item[2])
            parts.append(f'<div class="msg agent"><div class="who">Contoso assistant</div>{item[1]}'
                         f'{f"<div class=cites>{cites}</div>" if cites else ""}</div>')
        elif kind == "tool":
            _, name, arguments, running = item
            parts.append(f'<div class="tool{" running" if running else ""}">{name}<span>({esc(arguments)})</span>'
                         '</div>')
    if chips:
        parts.append('<div class="chips"><span class="chip">Track my order</span><span class="chip">Start a return'
                     '</span><span class="chip">Which tent fits four?</span></div>')
    composer = (f'<div class="composer typed">{esc(draft)}<span class="caret"></span><span class="send">↑</span></div>'
                if draft else '<div class="composer">Ask about an order, a return or our gear…<span class="send">↑'
                              '</span></div>')
    return f"""<div class="app"><div class="top"><div class="mark"></div><div class="brand">Contoso Outdoor<small>
Customer support</small></div><span class="pill">AI assistant</span><span class="note">Answers are AI-generated.
Check important details.</span></div><div class="thread">{"".join(parts)}</div>{composer}</div>"""


def browser_window(content: str, size: tuple[int, int]) -> str:
    toolbar = '<div class="nav"><b>‹</b><b>›</b></div><div class="url"><i></i>localhost:5173</div>'
    return window(content, "", size, False, CHAT_CSS, toolbar)


def tool(name: str, arguments: str, running: bool = False) -> tuple:
    return ("tool", name, arguments, running)


# The calls prompts/returns-assistant.md asks for: open the return as "undecided", then record the customer's choice.
LOOKUP = 'order_id="48213"'
OPENED = 'order_id="48213", reason="damaged", resolution="undecided"'
CHOSEN = 'order_id="48213", reason="damaged", resolution="replacement"'
FIRST_TURN = [("user", ASK), tool("lookup_order", LOOKUP), tool("create_return", OPENED), ANSWER]
CHAT_STATES = [
    ([GREETING], "", True, 1.8),
    ([GREETING], ASK[:31], True, 0.7),
    ([GREETING], ASK[:63], True, 0.7),
    ([GREETING], ASK, True, 0.9),
    ([GREETING, ("user", ASK), tool("lookup_order", LOOKUP, True)], "", False, 1.2),
    ([GREETING, ("user", ASK), tool("lookup_order", LOOKUP), tool("create_return", OPENED, True)], "", False, 1.0),
    ([GREETING, *FIRST_TURN], "", False, 3.4),
    (FIRST_TURN, ASK_AGAIN[:25], False, 0.7),
    (FIRST_TURN, ASK_AGAIN, False, 0.9),
    ([*FIRST_TURN, ("user", ASK_AGAIN), tool("create_return", CHOSEN, True)], "", False, 1.2),
    ([*FIRST_TURN[1:], ("user", ASK_AGAIN), tool("create_return", CHOSEN), DONE], "", False, 3.6),
]


@builder
def screenshots():
    """Screenshots/: azd deploying the support agent, and its chat answering. Named as macOS names them, which is
    how Biblo tells screenshots apart once the files have lost their screenshot attribute."""
    terminal = window(f"<pre>{TERMINAL}</pre>", "contoso-support-agent — azd up — 124×37", (1080, 770), True,
                      TERMINAL_CSS)
    save_png(render(terminal, (1192, 882), 2, True), out("Screenshots/Screenshot 2026-10-08 at 16.42.07.png"))
    thread, draft, chips, _ = CHAT_STATES[-1]
    page_ = browser_window(chat(thread, draft, chips), (1180, 800))
    save_png(render(page_, (1292, 912), 2, True), out("Screenshots/Screenshot 2026-10-09 at 12.44.10.png"))


EDITOR_CSS = """.content{background:#282c34;display:grid;grid-template-columns:48px 240px 1fr;
grid-template-rows:minmax(0,1fr) 24px;font:13px system-ui;color:#abb2bf}
.acts{border-right:1px solid #181a1f}.acts div{height:48px;display:grid;place-items:center}
.acts div.on{box-shadow:inset 2px 0 #d7dae0}
.acts svg{width:24px;height:24px;fill:none;stroke:#5c6370;stroke-width:1.6;stroke-linecap:round;stroke-linejoin:round}
.acts .on svg{stroke:#d7dae0}
.side{background:#21252b;border-right:1px solid #181a1f;padding-top:12px;overflow:hidden}
.side h4{margin:0 0 8px 20px;font:500 11px system-ui;letter-spacing:.06em;color:#9da5b4}
.row{height:22px;display:flex;align-items:center;white-space:nowrap}
.row.root{font:700 11px system-ui;letter-spacing:.04em;color:#d7dae0}
.row.on{background:#2c313a;color:#fff}.row.dim{color:#5c6370}
.row svg{width:16px;height:16px;flex:none;fill:none;stroke:#7f848e;stroke-width:1.4}
.row s,.tab s{text-decoration:none;font:700 8.5px "SF Mono",Menlo,monospace;width:20px;margin-right:4px;flex:none;
text-align:center}
.main{display:grid;grid-template-rows:35px 24px minmax(0,1fr) 190px;min-width:0}
.tabs{background:#21252b;display:flex}
.tab{padding:0 12px;display:flex;align-items:center;border-right:1px solid #181a1f;color:#7f848e}
.tab.on{background:#282c34;color:#e6e6e6}.tab .x{margin-left:12px;font-size:15px;color:#9da5b4}
.crumbs{padding-left:20px;display:flex;align-items:center;gap:9px;color:#7f848e;font-size:12.5px}
.crumbs .fn{color:#c678dd}
.editor{position:relative;overflow:hidden}
.ln{display:flex;height:18px;font:12px/18px "SF Mono",Menlo,monospace;white-space:pre}
.ln:first-child{margin-top:2px}
.ln b{width:62px;padding-right:24px;text-align:right;font-weight:400;color:#495162;flex:none}
.ln .t{position:relative;flex:1}.ln i{position:absolute;top:0;bottom:0;width:1px;background:#3b4048}
.ln.on{background:#2c313c}.ln.on b{color:#abb2bf}
.mini{position:absolute;right:16px;top:2px;opacity:.55}
.term{border-top:1px solid #181a1f;background:#21252b;padding:0 18px;overflow:hidden}
.term h5{margin:0;height:35px;display:flex;align-items:center;gap:24px;font:500 11px system-ui;letter-spacing:.06em;
color:#7f848e}.term h5 b{font-weight:500;color:#d7dae0;box-shadow:inset 0 -1px #d7dae0;padding:11px 0}
.term pre{margin:4px 0 0;font:12px/18px "SF Mono",Menlo,monospace;color:#abb2bf;white-space:pre}
.term .p{color:#61afef;font-weight:700}.term .m{color:#c678dd}.term .g{color:#98c379}.term .cursor{background:#abb2bf}
.status{grid-column:1/-1;background:#21252b;border-top:1px solid #181a1f;display:flex;align-items:center;gap:16px;
padding:0 12px;font-size:12px;color:#9da5b4}
.status svg{width:14px;height:14px;fill:none;stroke:currentColor;stroke-width:1.4;vertical-align:-3px;margin-right:4px}
.status .end{margin-left:auto;display:flex;gap:18px}"""

# The support agent's folder as the editor's explorer shows it: folders first, then files, each alphabetically, and
# what git ignores dimmed. (depth, open folder?, name, file type badge, ignored?)
AGENT_TREE = [(0, False, ".azure", "", True), (0, False, ".venv", "", True), (0, False, "evals", "", False),
              (0, False, "infra", "", False), (0, True, "prompts", "", False),
              (1, None, "returns-assistant.md", "md", False), (0, True, "src", "", False),
              (1, False, "__pycache__", "", True), (1, None, "__init__.py", "py", False),
              (1, None, "agent.py", "py", False), (1, None, "app.py", "py", False),
              (1, None, "evaluate.py", "py", False), (1, None, "tools.py", "py", False),
              (0, False, "tests", "", False), (0, False, "web", "", False), (0, None, ".gitignore", "git", False),
              (0, None, "azure.yaml", "yml", False), (0, None, "Dockerfile", "dk", False),
              (0, None, "pyproject.toml", "tml", False), (0, None, "README.md", "md", False)]
BADGES = {"py": "#e5c07b", "md": "#61afef", "git": "#e06c75", "yml": "#c678dd", "dk": "#61afef", "tml": "#9da5b4"}
CHEVRON = {True: '<svg viewBox="0 0 16 16"><path d="M4.5 6.5l3.5 3.5 3.5-3.5"/></svg>',
           False: '<svg viewBox="0 0 16 16"><path d="M6.5 4.5l3.5 3.5-3.5 3.5"/></svg>', None: '<svg></svg>'}


def colored(source: str) -> list[list[tuple[str, str, bool]]]:
    """Python as the One Dark theme colors it, a list of (text, color, italic) per line. The whole file at once, so
    a line in the middle of a docstring is still a string."""
    theme = get_style_by_name("one-dark")
    rows: list[list[tuple[str, str, bool]]] = [[]]
    for kind, value in lex(source, PythonLexer(stripnl=False)):
        look = theme.style_for_token(kind)
        for index, piece in enumerate(value.split("\n")):
            if index:
                rows.append([])
            if piece:
                rows[-1].append((piece, look["color"] or "abb2bf", look["italic"]))
    return rows


def indent_levels(lines: list[str]) -> list[int]:
    """Indent guides per line. A blank line keeps the guides the lines around it share."""
    levels = [(len(line) - len(line.lstrip())) // 4 if line.strip() else -1 for line in lines]
    for index, level in enumerate(levels):
        if level < 0:
            before = next((v for v in reversed(levels[:index]) if v >= 0), 0)
            after = next((v for v in levels[index + 1:] if v >= 0), 0)
            levels[index] = min(before, after)
    return levels


def minimap(rows: list[list[tuple[str, str, bool]]], first: int, count: int) -> str:
    """The editor's overview of the file: every token a sliver of its color, two pixels a line, and a slider
    over the lines on screen (`count` of them, from line `first`)."""
    rects = [f'<rect y="{2 * (first - 1)}" width="100" height="{2 * count}" fill="#fff" fill-opacity=".2"/>']
    for y, row in enumerate(rows):
        x = 0
        for piece, color, _ in row:
            if piece.strip():
                lead = len(piece) - len(piece.lstrip())
                rects.append(f'<rect x="{(x + lead) * .8:.1f}" y="{2 * y}" width="{len(piece.strip()) * .8:.1f}" '
                             f'height="1.5" fill="#{color}"/>')
            x += len(piece)
    return f'<svg class="mini" width="100" height="{2 * len(rows)}">{"".join(rects)}</svg>'


@builder
def editor():
    """Screenshots/: Sara, in the support agent's code, putting the current policies in a new vector store and
    publishing version 8. Her task from the meeting of October 6 (the agent had quoted the 2025 return policy), done
    between azd up (16:42) and the README (17:03). Words and shapes only, no product logos."""
    source = (OUT / "Code/contoso-support-agent/src/agent.py").read_text(encoding="utf-8")
    lines = source.splitlines()
    rows = colored(source)
    levels = indent_levels(lines)
    first = next(i for i, line in enumerate(lines, start=1) if line.startswith("def upload_policies")) - 4
    current = next(i for i, line in enumerate(lines, start=1) if "vector_stores.create(" in line)
    code = []
    for number in range(first, first + 32):
        text = "".join(f'<span style="color:#{color}{";font-style:italic" if italic else ""}">{esc(piece)}</span>'
                       for piece, color, italic in rows[number - 1])
        guides = "".join(f'<i style="left:{4 * k}ch"></i>' for k in range(levels[number - 1]))
        code.append(f'<div class="ln{" on" if number == current else ""}"><b>{number}</b>'
                    f'<span class="t">{guides}{text}</span></div>')
    acts = ['<path d="M13.5 3H6.5v18h11V7zM13.5 3v4h4"/>',
            '<circle cx="10.5" cy="10.5" r="6"/><path d="M15 15l5.5 5.5"/>',
            '<circle cx="6" cy="5" r="2.2"/><circle cx="6" cy="19" r="2.2"/><circle cx="18" cy="8" r="2.2"/>'
            '<path d="M6 7.2v9.6M18 10.2c0 4-5 4-11 7"/>', '<path d="M7 4.5l12 7.5-12 7.5z"/>',
            '<rect x="3.5" y="3.5" width="7" height="7"/><rect x="13.5" y="3.5" width="7" height="7"/>'
            '<rect x="3.5" y="13.5" width="7" height="7"/><rect x="13.5" y="13.5" width="7" height="7"/>']
    acts_html = "".join(f'<div{" class=on" if i == 0 else ""}><svg viewBox="0 0 24 24">{icon}</svg></div>'
                        for i, icon in enumerate(acts))
    tree = "".join(
        f'<div class="row{" on" if name == "agent.py" else ""}{" dim" if ignored else ""}" '
        f'style="padding-left:{10 + 12 * depth}px">{CHEVRON[open_]}'
        f'{f"<s style=color:{BADGES[kind]}>{kind}</s>" if kind else ""}{esc(name)}</div>'
        for depth, open_, name, kind, ignored in AGENT_TREE)
    store = "vs_7Qe2kR9mXa4LtW1nBc8ZpD5h"
    prompt = ('<span class="p">~/src/contoso-support-agent</span> <span class="m">main</span> '
              '<span class="g">❯</span> ')
    terminal = (f"{prompt}uv run python -m src.agent policies ~/Downloads/policies/*.pdf\n"
                f"9 pages in {store}. Next: azd env set POLICY_VECTOR_STORE_ID {store}, then publish.\n"
                f"{prompt}azd env set POLICY_VECTOR_STORE_ID {store}\n"
                f'{prompt}set -a; eval "$(azd env get-values)"; set +a\n'
                f"{prompt}uv run python -m src.agent publish\n"
                f"Published returns-assistant version 8. Next: python -m src.evaluate --version 8\n"
                f'{prompt}<span class="cursor"> </span>')
    branch = ('<svg viewBox="0 0 16 16"><circle cx="4.5" cy="3.5" r="1.6"/><circle cx="4.5" cy="12.5" r="1.6"/>'
              '<circle cx="11.5" cy="5" r="1.6"/><path d="M4.5 5.1v5.8M11.5 6.6c0 2.6-3.5 2.6-7 4.6"/></svg>')
    errors = ('<svg viewBox="0 0 16 16"><circle cx="8" cy="8" r="6"/>'
              '<path d="M5.8 5.8l4.4 4.4M10.2 5.8l-4.4 4.4"/></svg>')
    warnings = '<svg viewBox="0 0 16 16"><path d="M8 2.2l6.3 11.3H1.7zM8 6.5v3.4M8 11.6v.3"/></svg>'
    body = f"""<div class="acts">{acts_html}</div>
<div class="side"><h4>EXPLORER</h4><div class="row root" style="padding-left:2px">{CHEVRON[True]}CONTOSO-SUPPORT-AGENT
</div>{tree}</div>
<div class="main"><div class="tabs"><div class="tab on"><s style="color:{BADGES['py']}">py</s>agent.py
<span class="x">×</span></div><div class="tab"><s style="color:{BADGES['py']}">py</s>tools.py</div>
<div class="tab"><s style="color:{BADGES['md']}">md</s>returns-assistant.md</div></div>
<div class="crumbs">src<span>›</span>agent.py<span>›</span><span class="fn">ƒ</span>upload_policies</div>
<div class="editor">{"".join(code)}{minimap(rows, first, len(code))}</div>
<div class="term"><h5><span>PROBLEMS</span><span>OUTPUT</span><span>DEBUG CONSOLE</span><b>TERMINAL</b>
<span>PORTS</span></h5><pre>{terminal}</pre></div></div>
<div class="status"><span>{branch}main</span><span>{errors}0&nbsp;&nbsp;{warnings}0</span><div class="end">
<span>Ln {current}, Col 5</span><span>Spaces: 4</span><span>UTF-8</span><span>LF</span><span>Python</span>
<span>3.12.7 ('.venv': venv)</span></div></div>"""
    shot = window(body, "agent.py — contoso-support-agent", (1440, 900), True, EDITOR_CSS)
    save_png(render(shot, (1552, 1012), 2, True), out("Screenshots/Screenshot 2026-10-08 at 16.55.18.png"))


# MARK: Dates

# Files whose contents give them a day. Screenshots and screen recordings take theirs from their names; everything else
# gets a steady pseudo-random working hour between FIRST and LAST.
DATES = {
    "Documents/Notas de reunión – piloto de agentes (Contoso).docx": "2026-10-06 17:52",
    "Audio/Resumen semanal – piloto de agentes.m4a": "2026-10-09 18:20",
    "Images/IMG_4127.jpg": "2026-09-29 11:06",
    "Images/Copilot adoption by team – Q3 2026.png": "2026-10-02 09:41",
    "Slides/Microsoft Fabric – platform overview.pptx": "2026-10-07 18:40",
    "Slides/Taller práctico – agentes con Microsoft Foundry.pptx": "2026-10-09 13:15",
    "Code/contoso-support-agent/infra/resources.bicep": "2026-10-08 15:58",
    "Code/contoso-support-agent/README.md": "2026-10-08 17:03",
    "Code/contoso-support-agent/src/agent.py": "2026-10-08 11:27",
    "Code/contoso-support-agent/evals/sample.jsonl": "2026-10-09 10:31",
    "Code/contoso-support-agent/src/evaluate.py": "2026-10-09 10:48",
    "Code/fabric-notebooks/add_order_history_shortcut.py": "2026-10-09 09:12",
}
# Projects written over part of the pilot, not all of it.
SPANS = {
    "Code/contoso-support-agent/": (dt.datetime(2026, 10, 1), dt.datetime(2026, 10, 8)),
    "Code/fabric-notebooks/": (dt.datetime(2026, 9, 28), dt.datetime(2026, 10, 7)),
}
FIRST, LAST = dt.datetime(2026, 9, 21), dt.datetime(2026, 10, 11)
NAMED = re.compile(r"(?:Screenshot|Screen Recording) (\d{4})-(\d\d)-(\d\d) at (\d\d)\.(\d\d)\.(\d\d)")


def working_hour(key: str, first: dt.datetime, last: dt.datetime) -> dt.datetime:
    rng = random.Random(key)
    day = first + dt.timedelta(days=rng.randrange((last - first).days + 1))
    return day.replace(hour=rng.randint(9, 18), minute=rng.randrange(60), second=rng.randrange(60))


def stamp() -> int:
    count = 0
    for path in sorted(OUT.rglob("*")):
        if not path.is_file() or path.name == ".DS_Store":
            continue
        relative = path.relative_to(OUT).as_posix()
        if match := NAMED.search(path.name):
            when = dt.datetime(*map(int, match.groups()))
        elif relative in DATES:
            when = dt.datetime.strptime(DATES[relative], "%Y-%m-%d %H:%M")
        else:
            span = next((s for prefix, s in SPANS.items() if relative.startswith(prefix)), (FIRST, LAST))
            when = working_hour(relative, *span)
        os.utime(path, (when.timestamp(), when.timestamp()))
        count += 1
    return count


def main() -> None:
    parser = argparse.ArgumentParser(description="Makes Biblo's demo files in Resources/Demo.noindex.",
                                     epilog="groups: " + ", ".join(BUILDERS))
    parser.add_argument("only", nargs="*", help="groups to make again (default: all of them, from scratch)")
    arguments = parser.parse_args()
    if unknown := [name for name in arguments.only if name not in BUILDERS]:
        parser.error(f"no group called {', '.join(unknown)}; the groups are {', '.join(BUILDERS)}")
    global BROWSER
    BROWSER = find_browser()
    OUT.mkdir(parents=True, exist_ok=True)
    if not arguments.only:
        for path in OUT.rglob("*"):
            if path.is_file() and path.suffix.lower() in MADE:
                path.unlink()
    try:
        for name in arguments.only or list(BUILDERS):
            started = time.monotonic()
            BUILDERS[name]()
            print(f"{name}: {time.monotonic() - started:.0f} s", flush=True)
        files = stamp()
    finally:
        shutil.rmtree(WORK, ignore_errors=True)
    size = sum(path.stat().st_size for path in OUT.rglob("*") if path.is_file())
    print(f"{OUT.relative_to(HERE.parents[1])}: {files} files, {size / 1e6:.1f} MB")



@builder
def screencast():
    """Video/: the support chat answering, recorded as ⇧⌘5 records a part of the screen (no sound)."""
    size, scale, margin = (1180, 800), 1.5, 56
    box = [int(v * scale) for v in (margin, margin - 16, margin + size[0], margin - 16 + size[1])]
    states = []
    for thread, draft, chips, hold in CHAT_STATES:
        shot = render(browser_window(chat(thread, draft, chips), size), (size[0] + 112, size[1] + 112), scale, True)
        window_only = shot.crop(box)
        frame = Image.new("RGB", window_only.size, (228, 228, 231))
        frame.paste(window_only, (0, 0), window_only)
        states.append((frame, hold))
    recording(states, out("Video/Screen Recording 2026-10-09 at 12.41.33.mov"))


# Narrated explainers, made only of words and shapes: no logos in videos (assets/SOURCES.md).
# Each slide: kicker, title, line under it, chips, narration.
FOUNDRY_VIDEO = [
    ("Explainer", "Microsoft Foundry in 60 seconds", "", [], "Microsoft Foundry, in sixty seconds."),
    ("What it is", "One platform for AI apps and agents.",
     "Models, agents, tools and monitoring, together in one Azure resource.", [],
     "Foundry is Microsoft's platform for AI apps and agents: models, agents, tools and monitoring, together in "
     "one Azure resource."),
    ("Models", "Pick a model. Deploy it in minutes.",
     "Compare them on your own prompts, or let the model router choose.",
     ["OpenAI", "Anthropic", "Meta", "Mistral", "DeepSeek", "xAI", "Microsoft"],
     "Start in the model catalog: OpenAI, Anthropic, Meta, Mistral, DeepSeek, xAI, Microsoft. Compare them on "
     "your own prompts, or let the model router choose."),
    ("Agents", "Agents that act, not just answer.", "Foundry Agent Service runs them: instructions, a model, tools.",
     ["File search", "Code interpreter", "Web search", "MCP servers"],
     "Foundry Agent Service runs your agents: instructions, a model, and tools like file search, code interpreter "
     "or any MCP server."),
    ("Knowledge", "Grounded in your own data.", "Answers cite the documents they come from.",
     ["Your files", "Azure AI Search", "SharePoint", "Microsoft Fabric"],
     "Ground the answers in your own files, Azure AI Search, SharePoint or Microsoft Fabric, with citations."),
    ("Observe and evaluate", "See every step. Measure every release.", "",
     ["Tracing", "Evaluations", "Monitoring", "AI red teaming"],
     "Every run is traced, and built-in evaluators score groundedness, relevance and tool-call accuracy."),
    ("Trust", "Safe by default.", "", ["Guardrails", "Microsoft Entra identities", "Private networking"],
     "Guardrails check prompts, responses and tool calls, and agents sign in with Entra identities, not keys."),
    ("Try it", "ai.azure.com", "Sign in with your Azure account and create a project.", [],
     "Try it today, at ai dot azure dot com."),
]
ONELAKE_VIDEO = [
    ("Microsoft Fabric", "OneLake, explicado en 30 segundos", "", [], "OneLake, explicado en treinta segundos."),
    ("Qué es", "Un solo lago de datos para toda la organización.", "Como un OneDrive, pero para los datos.", [],
     "OneLake es el lago de datos de Microsoft Fabric: uno para toda la organización, como un OneDrive para los "
     "datos."),
    ("Una sola copia", "Una copia. Todos los motores.", "Tablas en formato abierto Delta Parquet.",
     ["SQL", "Spark", "Power BI", "Real-Time Intelligence"],
     "Las tablas se guardan en formato abierto Delta Parquet, y todos los motores leen la misma copia."),
    ("Accesos directos", "Conecta datos sin moverlos.", "",
     ["Azure Data Lake", "Amazon S3", "Google Cloud Storage", "Dataverse"],
     "Con los accesos directos conectas datos de Azure, Amazon S3 o Google Cloud, sin copiarlos."),
    ("Gobierno", "Encuentra y protege tus datos en un solo sitio.", "Con el catálogo de OneLake.", [],
     "Y en el catálogo de OneLake encuentras, gobiernas y proteges todos tus datos."),
]

EXPLAINER_CSS = """.text{position:absolute;top:0;bottom:0;display:flex;flex-direction:column;justify-content:center}
.kicker{font:600 15px/1 "SF Mono",Menlo,monospace;letter-spacing:.16em;text-transform:uppercase;margin-bottom:26px}
h1{margin:0;text-wrap:balance}
p{font:400 25px/1.4 system-ui;margin:24px 0 0;text-wrap:pretty}
.chips{display:flex;flex-wrap:wrap;gap:8px;margin-top:34px}
.chip{padding:9px 14px;border-radius:999px;font:500 18px/1.2 system-ui}
.foot{position:absolute;bottom:42px;display:flex;justify-content:space-between;
font:500 13px "SF Mono",Menlo,monospace;letter-spacing:.1em;text-transform:uppercase}
"""
FOUNDRY_CSS = EXPLAINER_CSS + """body{background:#0a0e24;color:#eef2ff}
.glow{position:absolute;inset:0;background:
radial-gradient(760px 560px at 79.5% 50%,rgba(99,102,241,.30),transparent 62%),
radial-gradient(640px 420px at 0% 110%,rgba(56,189,248,.10),transparent 60%)}
.text{left:96px;width:700px}.kicker{color:#a5b4fc}
h1{font:700 60px/1.05 system-ui;letter-spacing:-.025em;color:#fff}p{color:#c7d2fe;max-width:640px}
.chip{border:1px solid rgba(165,180,252,.34);background:rgba(99,102,241,.13);color:#e0e7ff}
.dial{position:absolute;right:92px;top:50%;width:340px;height:340px;transform:translateY(-50%)}
.foot{left:96px;right:96px;color:#6870a6}"""
ONELAKE_CSS = EXPLAINER_CSS + """body{background:#f1f6f4;color:#0b3b36}
.text{left:104px;width:680px}.kicker{color:#0d9488}
h1{font:600 60px/1.08 "Iowan Old Style",Charter,Georgia,serif;letter-spacing:-.01em;color:#0b3b36}
p{color:#3f6b65}
.chip{border:1px solid #9fd8ce;background:#fff;color:#0b3b36}
.lake{position:absolute;left:800px;top:240px;width:600px;height:600px}
.foot{left:104px;justify-content:flex-start;gap:28px;color:#6d958f}"""


def stopwatch(elapsed: float) -> str:
    """The Foundry explainer's dial: how far into its minute each slide starts."""
    c, r = 180, 148
    ticks = "".join(f'<line x1="{c}" y1="{c - r - 12}" x2="{c}" y2="{c - r - (26 if i % 5 == 0 else 18)}" '
                    f'transform="rotate({i * 6} {c} {c})"/>' for i in range(60))
    arc = 2 * math.pi * r * min(elapsed, 60) / 60
    second = int(elapsed)
    return f"""<svg class="dial" viewBox="0 0 360 360"><g stroke="rgba(165,180,252,.38)" stroke-width="2">{ticks}</g>
<circle cx="{c}" cy="{c}" r="{r}" fill="none" stroke="rgba(165,180,252,.14)" stroke-width="2"/>
<circle cx="{c}" cy="{c}" r="{r}" fill="none" stroke="#818cf8" stroke-width="7" stroke-linecap="round"
stroke-dasharray="{arc:.1f} 2000" transform="rotate(-90 {c} {c})"/>
<text x="{c}" y="{c + 24}" text-anchor="middle" fill="#fff" font-family="system-ui" font-size="70"
font-weight="200">{second // 60}:{second % 60:02d}</text></svg>"""


def ripples(count: int) -> str:
    """The OneLake explainer's lake: one more ring for every slide."""
    rings = "".join(f'<circle r="{48 + 60 * i}" fill="none" stroke="#0d9488" '
                    f'stroke-width="{2.5 if i == count - 1 else 1.5}" '
                    f'opacity="{1 if i == count - 1 else round(0.5 - 0.07 * i, 2)}"/>' for i in range(count))
    return f'<svg class="lake" viewBox="-300 -300 600 600"><circle r="14" fill="#0d9488"/>{rings}</svg>'


def explainer(script: list, css: str, motif, foot: str, speaker: tuple[str, str], rate: int, target: float,
              path: Path) -> None:
    """A narrated explainer, about `target` seconds long: every slide stays on while its line is spoken."""
    voices = [say(slide[-1], voice(*speaker), rate) for slide in script]
    spoken = [seconds(v) for v in voices]
    fade = 0.5
    pad = min(max((target - fade - sum(spoken)) / len(script), 0.45), 1.6)
    images, start = [], 0.0
    for index, (kicker, title, line, chips, _) in enumerate(script):
        tags = "".join(f'<span class="chip">{esc(chip)}</span>' for chip in chips)
        body = (f'{motif(index, start)}<div class="text"><div class="kicker">{esc(kicker)}'
                f'</div><h1>{esc(title)}</h1>{f"<p>{esc(line)}</p>" if line else ""}'
                f'{f"<div class=chips>{tags}</div>" if tags else ""}</div>'
                f'<div class="foot"><span>{esc(foot)}</span><span>{index + 1:02d} / {len(script):02d}</span></div>')
        images.append(render(page(body, css), (1280, 720), 1.5))
        start += spoken[index] + pad
    slideshow(list(zip(images, voices)), path, fade=fade, pad=pad)


@builder
def videos():
    """Video/: two narrated explainers, Microsoft Foundry in English and OneLake in Spanish."""
    explainer(FOUNDRY_VIDEO, FOUNDRY_CSS, lambda index, start: f'<div class="glow"></div>{stopwatch(start)}',
              "Microsoft Foundry · explainer",
              ("Moira", "en_IE"), 185, 60, out("Video/Foundry in 60 seconds.mp4"))
    explainer(ONELAKE_VIDEO, ONELAKE_CSS, lambda index, start: ripples(index + 1), "OneLake · Microsoft Fabric",
              ("Paulina", "es_MX"), 190, 30, out("Video/OneLake explicado en 30 segundos.mp4"))


# MARK: Audio

_voices: list[tuple[str, str]] = []


def voice(name: str, locale: str) -> str:
    """A macOS voice by its name and locale: their full names are in the Mac's language ("Reed (Spanish (Spain))")."""
    if not _voices:
        for line in run("say", "-v", "?").splitlines():
            if match := re.match(r"(.+?)\s+([a-z]{2}_[A-Z]{2})\s+#", line):
                _voices.append((match.group(1).strip(), match.group(2)))
    for full, where in _voices:
        if where == locale and (full == name or full.startswith(name + " (")):
            return full
    sys.exit(f"needs the macOS voice {name} ({locale}): System Settings → Accessibility → Spoken Content → "
             "System voice → Manage Voices…")


def conversation(lines: list[tuple[str, str]], voices: dict[str, tuple[str, str]], path: Path,
                 pause: float = 0.45) -> None:
    parts = [(say(text, voice(*voices[who])), pause) for who, text in lines]
    join_audio(parts, path)


@builder
def audio():
    """Audio/: a podcast episode, voice memos and a weekly summary, in English and Spanish."""
    conversation([
        ("S", "Welcome back to AI in Two Minutes. I'm Sam."),
        ("D", "And I'm Daniel. Today's question: what actually is an AI agent?"),
        ("S", "Short version: an agent is a language model that can take actions. You give it instructions, a goal, "
              "and a set of tools, and it decides which steps to take."),
        ("D", "So a chatbot answers. An agent does things."),
        ("S", "Exactly. Say a customer writes: my tent arrived broken. The agent can look up the order, check the "
              "return policy, open a return, and reply, all in one conversation."),
        ("D", "And the tools are just functions, or APIs?"),
        ("S", "Right. An order lookup, a search over your documents, a calendar, running code. The model picks the "
              "tool and fills in the arguments; your code runs it and hands back the result."),
        ("D", "What's the catch?"),
        ("S", "Agents make mistakes too. So you ground them in your own data, limit what each tool is allowed to do, "
              "keep a human in the loop for anything risky, and evaluate them before launch and while they run."),
        ("D", "Ground, limit, evaluate. That's our two minutes. See you next week!"),
    ], {"S": ("Samantha", "en_US"), "D": ("Daniel", "en_GB")},
        out("Audio/AI in Two Minutes – Ep. 12 – What is an AI agent.m4a"))

    conversation([("M", "Nota rápida para el jueves. En la demo de Microsoft Fabric para el equipo de ventas quiero "
                        "enseñar tres cosas. Primero, OneLake: que los pedidos y las devoluciones están en un único "
                        "lago de datos, sin copias. Segundo, un informe de Power BI en modo Direct Lake, que lee las "
                        "tablas directamente, sin importarlas. Y tercero, Copilot en Fabric: pedirle en lenguaje "
                        "natural un resumen de ventas por región. Ojo, usar solo los datos de ejemplo de Contoso, "
                        "nada de datos reales de clientes. Y llevar el portátil cargado.")],
                 {"M": ("Mónica", "es_ES")}, out("Audio/Nota de voz – demo de Fabric.m4a"))

    conversation([("K", "Reminder for Monday's deck. On the Copilot slide, use the official logo files only: "
                        "download them from the brand pages, don't screenshot them from a website. No recoloring, "
                        "no stretching, no drop shadows, and don't animate them. Leave clear space around the logo. "
                        "The Microsoft Copilot icon has a new design, so check with the brand team which version to "
                        "use before anything goes public. And for GitHub Copilot, use the lockup with the wordmark, "
                        "not just the icon.")],
                 {"K": ("Karen", "en_AU")}, out("Audio/Voice memo – Copilot logos for the deck.m4a"))

    conversation([
        ("M", "Resumen de la semana del piloto del asistente de devoluciones."),
        ("R", "Seguimos con cincuenta usuarios del equipo de soporte. Esta semana el agente resolvió el sesenta y dos "
              "por ciento de las consultas de devoluciones sin pasarlas a una persona."),
        ("M", "¿Y la satisfacción?"),
        ("R", "Cuatro coma cinco sobre cinco en las encuestas. Y el tiempo medio de resolución bajó de once minutos a "
              "cuatro."),
        ("M", "¿Algún problema?"),
        ("R", "Dos. Con pedidos de hace más de un año, la herramienta de búsqueda de pedidos falla. Y en tres casos "
              "el agente citó una política de devoluciones antigua."),
        ("M", "Vale. Acciones: actualizar el índice de documentos, añadir esos casos al conjunto de evaluación, y "
              "repetir las evaluaciones antes de ampliar el piloto a ventas."),
        ("R", "Perfecto. Hablamos el viernes que viene."),
    ], {"M": ("Mónica", "es_ES"), "R": ("Reed", "es_ES")}, out("Audio/Resumen semanal – piloto de agentes.m4a"))


# MARK: Documents

LOGO_FACTS_CSS = """@page{size:A4;margin:15mm 17mm 14mm}
html,body{width:auto;height:auto;overflow:visible}
body{font-size:10pt;line-height:1.45;color:#1d1d1f}
.kicker{font:600 8pt system-ui;letter-spacing:.12em;text-transform:uppercase;color:#6e6e73}
h1{font-size:27pt;letter-spacing:-.02em;margin:4px 0 6px;line-height:1.1}
.lede{font-size:11.5pt;color:#3a3a3c;margin:0 0 10px;max-width:150mm}
h2{font-size:14pt;margin:14px 0 6px;padding-top:10px;border-top:1px solid #e5e5ea}
.page2 h2:first-child{border-top:0;padding-top:0;margin-top:0}
.row{display:grid;grid-template-columns:1fr 54mm;gap:9mm;align-items:start}
ol.time{list-style:none;margin:0;padding:0}
ol.time li{display:grid;grid-template-columns:27mm 1fr;gap:3mm;padding:3px 0}
ol.time b{font-variant-numeric:tabular-nums;color:#0f6cbd;font-weight:600}
figure{margin:4px 0 0;padding:12px;border:1px solid #e5e5ea;border-radius:10px;background:#fff}
figure .art{height:30mm;display:grid;place-items:center;background:#fff}
figure img{height:auto}
figcaption{font-size:8.4pt;color:#6e6e73;margin-top:8px;line-height:1.4}
.tag{display:inline-block;font:600 7.4pt system-ui;letter-spacing:.06em;text-transform:uppercase;padding:2px 6px;
border-radius:4px;background:#fff4ce;color:#7a5900;margin-bottom:4px}
.rules{display:grid;grid-template-columns:1fr 1fr;gap:6mm;margin-top:4px;break-inside:avoid}
.rules div{border-radius:10px;padding:10px 14px}
.do{background:#eef8f1}.dont{background:#fdeeee}
.rules h3{margin:0 0 4px;font-size:10.5pt}.do h3{color:#1f7a45}.dont h3{color:#b42318}
.rules ul{margin:0;padding-left:16px}.rules li{margin:3px 0}
blockquote{margin:10px 0;padding:2px 0 2px 12px;border-left:3px solid #d1d1d6;color:#3a3a3c;font-size:9.8pt}
blockquote cite{display:block;font-style:normal;color:#6e6e73;font-size:8.4pt}
.sources{font-size:8.2pt;color:#6e6e73;columns:2;column-gap:8mm}
.sources p{margin:0 0 5px;break-inside:avoid}
footer{margin-top:12px;font-size:8pt;color:#8e8e93}
.page2{break-before:page}"""

LOGO_FACTS = """<div class="kicker">Contoso brand studio · one-pager · October 2026</div>
<h1>Copilot logos – quick facts</h1>
<p class="lede">There are two products called Copilot, with two different identities: <b>Microsoft Copilot</b> and
<b>GitHub Copilot</b>. Use the right one, use the official files, and check the current guidelines before anything goes
public.</p>

<h2>Microsoft Copilot</h2>
<div class="row"><ol class="time">
<li><b>Sep 21, 2023</b><span>Microsoft announces “Microsoft Copilot, your everyday AI companion”. It starts rolling out
in Windows 11 on September 26.</span></li>
<li><b>Nov 15, 2023</b><span>Bing Chat and Bing Chat Enterprise become Copilot, and copilot.microsoft.com
launches.</span></li>
<li><b>Oct 1, 2024</b><span>The consumer app gets a refreshed, warmer design.</span></li>
<li><b>Jan 15, 2025</b><span>The Microsoft 365 app becomes the Microsoft 365 Copilot app, with a new name and
icon.</span></li>
<li><b>2026</b><span>The Microsoft 365 (Office) app becomes the <b>Microsoft Copilot app</b>, with “a simpler name and
icon”. Web addresses move to copilot.cloud.microsoft.</span></li>
<li><b>Oct 5, 2026</b><span>Microsoft Design explains the new mark in <i>More than a mark</i>. It starts from the square:
“two distinct forms that make one silhouette”. The wordmark pairs Segoe with Segoe Serif, the color “keeps a thread back
to Copilot's original gradient”, and motion gestures stand for thinking, planning and retrieval.</span></li>
</ol>
<figure><div class="art"><img src="{copilot}" style="width:22mm" alt="Microsoft Copilot icon, 2023 design"></div>
<figcaption><span class="tag">Superseded</span><br>The Microsoft Copilot icon used from 2023: “a single, complex
overlapping ribbon” with a blue-to-pink gradient. Official file, shown as provided.</figcaption></figure></div>
<blockquote>As of October 12, 2026 we found no official download of the new Microsoft Copilot icon. Don't redraw it or
take it from a screenshot: ask the brand team which version to use.</blockquote>

<h2>GitHub Copilot</h2>
<div class="row"><ol class="time">
<li><b>Jun 29, 2021</b><span>Technical preview of GitHub Copilot, “powered by OpenAI Codex”.</span></li>
<li><b>Jun 6, 2022</b><span>The Copilot icon joins Primer Octicons, GitHub's icon set, at 16, 24, 48 and 96 px.</span></li>
<li><b>Jun 21, 2022</b><span>GitHub Copilot is generally available to individual developers.</span></li>
<li><b>Jun 7, 2023</b><span>The icon is redrawn “to tighten up Copilot branding”.</span></li>
<li><b>2025</b><span>“GitHub Copilot no longer has a standalone logo that heros the Copilot icon.” Use the <b>GitHub
Copilot lockup</b> instead. The Copilot palette “features purples, pinks, and oranges”.</span></li>
</ol>
<figure><div class="art"><img src="{lockup}" style="width:44mm" alt="GitHub Copilot lockup"></div>
<figcaption>The GitHub Copilot lockup: the GitHub mark with the words “GitHub Copilot”. Official file, black on white,
shown as provided.</figcaption></figure></div>

<div class="page2">
<h2>Do and don't</h2>
<div class="rules"><div class="do"><h3>Do</h3><ul>
<li>Download logos from the official brand pages.</li>
<li>Scale them evenly and keep the clear space around them.</li>
<li>Use the lockup when you mean GitHub Copilot.</li>
<li>Ask the brand owner before any public or commercial use.</li></ul></div>
<div class="dont"><h3>Don't</h3><ul>
<li>Recolor, stretch, skew, rotate or crop a logo.</li>
<li>Add shadows, glows or other effects.</li>
<li>Animate Microsoft logos.</li>
<li>Combine a logo with your own words or graphics, or use it as your own product's logo.</li>
<li>Take logos from screenshots or redraw them.</li></ul></div></div>

<blockquote>“Don't alter, animate, distort, or misappropriate Microsoft's Brand Assets.” Logos and icons “can never be used
without an express license”.<cite>Microsoft Trademark and Brand Guidelines</cite></blockquote>
<blockquote>“Do not modify the permitted GitHub logos, including changing the color, dimensions, or combining with other
words or design elements.”<cite>GitHub logo guidelines</cite></blockquote>
<blockquote>Azure architecture icons may be used “in architectural diagrams, training materials, or documentation”.
<cite>Azure Architecture Icons, Microsoft Learn</cite></blockquote>

<h2>Sources</h2>
<div class="sources">
<p>Announcing Microsoft Copilot, your everyday AI companion (Sep 21, 2023) –
blogs.microsoft.com/blog/2023/09/21/announcing-microsoft-copilot-your-everyday-ai-companion</p>
<p>Our vision to bring Microsoft Copilot to everyone (Nov 15, 2023) –
blogs.bing.com/search/november-2023/our-vision-to-bring-microsoft-copilot-to-everyone-and-more</p>
<p>An AI companion for everyone (Oct 1, 2024) – blogs.microsoft.com/blog/2024/10/01/an-ai-companion-for-everyone</p>
<p>The Microsoft 365 app transition to the Microsoft 365 Copilot app – support.microsoft.com</p>
<p>More than a mark (Oct 5, 2026) – microsoft.design/articles/more-than-a-mark</p>
<p>Introducing GitHub Copilot (Jun 29, 2021) and GitHub Copilot is generally available (Jun 21, 2022) – github.blog</p>
<p>primer/octicons pull requests 794 and 957 – github.com/primer/octicons</p>
<p>GitHub Copilot brand identity – brand.github.com/brand-identity/copilot</p>
<p>GitHub logo guidelines – brand.github.com/foundations/logo</p>
<p>Microsoft Trademark and Brand Guidelines – microsoft.com/en-us/legal/intellectualproperty/trademarks</p>
<p>Azure Architecture Icons – learn.microsoft.com/en-us/azure/architecture/icons</p>
</div>
<footer>Contoso is a fictitious company. This summary is not legal advice: the brand owners' current guidelines always
win.</footer></div>"""

MEETING_NOTES = """<h1>Notas de reunión – Piloto de agentes (Contoso)</h1>
<p><b>Fecha:</b> martes, 6 de octubre de 2026, de 16:30 a 17:30<br>
<b>Asistentes:</b> Lucía (producto, organiza), Javier (datos), Ana (soporte), Sara (desarrollo), Marc (seguridad)<br>
<b>Redacta:</b> Ana</p>
<h2>Orden del día</h2>
<ol><li>Estado del piloto del asistente de devoluciones (semana 2)</li>
<li>Calidad de las respuestas</li>
<li>Datos de pedidos en Microsoft Fabric</li>
<li>¿Ampliamos el piloto al equipo de ventas?</li></ol>
<h2>1. Estado del piloto</h2>
<ul><li>50 personas del equipo de soporte usan el asistente, que está construido como un agente de Microsoft Foundry
con dos herramientas: <i>lookup_order</i> y <i>create_return</i>.</li>
<li>Esta semana resolvió el 58 % de las consultas de devoluciones sin pasarlas a una persona. El objetivo es el
60 %.</li>
<li>Tiempo medio de resolución: 5 minutos (antes del piloto, 11). Satisfacción en la encuesta: 4,4 sobre 5.</li></ul>
<h2>2. Calidad de las respuestas</h2>
<ul><li>En tres conversaciones el agente citó la política de devoluciones de 2025: el índice de documentos no se
actualizó cuando cambió el PDF.</li>
<li>Con pedidos de hace más de un año, <i>lookup_order</i> no encuentra nada, porque el histórico no está en el
lakehouse.</li>
<li>Las evaluaciones automáticas lo confirman: <i>groundedness</i> 3,8 sobre 5, por debajo del umbral de 4.</li></ul>
<h2>3. Datos en Fabric</h2>
<ul><li>Javier propone crear un acceso directo (<i>shortcut</i>) de OneLake al histórico de pedidos en lugar de
copiarlo: los datos se quedan donde están y el lakehouse los lee como si fueran locales.</li>
<li>El jueves hay demo de Fabric para ventas: OneLake, un informe de Power BI en modo Direct Lake y Copilot en Fabric.
Solo con datos de ejemplo.</li></ul>
<h2>4. Ampliación a ventas</h2>
<p>No se amplía hasta repetir las evaluaciones con los casos nuevos y llegar a <i>groundedness</i> 4 o más.</p>
<h2>Decisiones</h2>
<ul><li>Reindexar la política de devoluciones vigente y quitar la antigua del índice.</li>
<li>Añadir al conjunto de evaluación los casos que fallaron esta semana.</li>
<li>Mantener las <i>guardrails</i> también en las llamadas a herramientas.</li></ul>
<h2>Próximos pasos</h2>
<table border="1" cellpadding="6" cellspacing="0">
<tr><th>Acción</th><th>Responsable</th><th>Fecha</th></tr>
<tr><td>Actualizar el índice de documentos</td><td>Sara</td><td>8 de octubre</td></tr>
<tr><td>Acceso directo de OneLake al histórico de pedidos</td><td>Javier</td><td>9 de octubre</td></tr>
<tr><td>Añadir 20 casos al conjunto de evaluación</td><td>Ana</td><td>9 de octubre</td></tr>
<tr><td>Revisar con Legal la retención de conversaciones</td><td>Marc</td><td>13 de octubre</td></tr>
<tr><td>Taller de agentes de Foundry para el equipo</td><td>Lucía</td><td>14 de octubre</td></tr>
</table>
<p><b>Próxima reunión:</b> martes, 13 de octubre.</p>
<p><i>Contoso es una empresa ficticia. Notas de ejemplo para una demo.</i></p>"""

GLOSSARY = [
    ("Acceso directo (shortcut)", "Objeto de OneLake que apunta a datos guardados en otro sitio: otra área de trabajo, "
     "Azure Data Lake Storage Gen2, Amazon S3 o Google Cloud Storage. Funciona como un enlace simbólico: los datos no "
     "se copian."),
    ("Área de trabajo (workspace)", "Contenedor de elementos de Fabric para un equipo o un proyecto, con sus propios "
     "permisos."),
    ("Capacidad (capacity)", "Recursos de proceso que compra la organización y que comparten todas las cargas de "
     "trabajo de Fabric."),
    ("Copilot en Fabric", "Ayuda a escribir consultas, pipelines y código, y a sacar resúmenes e ideas de los datos."),
    ("Data Factory", "Integración de datos: pipelines y flujos de datos (Dataflow Gen2) para traer y transformar "
     "datos."),
    ("Delta Lake", "Formato abierto de tablas, sobre archivos Parquet, en el que Fabric guarda las tablas de OneLake."),
    ("Direct Lake", "Modo de los modelos semánticos de Power BI que lee las tablas Delta de OneLake directamente, sin "
     "importarlas."),
    ("Eventhouse", "Base de datos de Real-Time Intelligence para eventos y series temporales. Se consulta con KQL."),
    ("Fabric IQ (versión preliminar)", "Lleva la capa semántica de Power BI a las operaciones del negocio: ontología, "
     "planificación, Fabric Graph, agentes de datos y de operaciones, y modelos semánticos."),
    ("KQL (Kusto Query Language)", "Lenguaje de consultas de Real-Time Intelligence."),
    ("Lakehouse", "Elemento que reúne archivos y tablas Delta en OneLake. Se trabaja con Spark y tiene un punto de "
     "conexión de análisis SQL."),
    ("Mirroring (creación de reflejos)", "Replica de forma continua una base de datos existente en OneLake: Azure SQL "
     "Database, Azure Cosmos DB, Snowflake, Google BigQuery…"),
    ("Modelo semántico", "La capa de negocio de Power BI: tablas, relaciones y medidas."),
    ("Notebook", "Cuaderno de Apache Spark (Python, Scala, SQL o R) para ingeniería y ciencia de datos."),
    ("OneLake", "El lago de datos único de toda la organización, construido sobre Azure Data Lake Storage Gen2. Una "
     "sola copia de los datos para todas las cargas de trabajo."),
    ("Pipeline", "Orquestación de Data Factory: copiar datos, transformarlos y programar cuándo se hace."),
    ("Power BI", "Informes y paneles. En Fabric es una carga de trabajo más sobre OneLake."),
    ("Real-Time Intelligence", "Datos en streaming y eventos: Eventstream, Eventhouse, Activator y el Real-Time hub."),
    ("SQL database en Fabric", "Base de datos transaccional basada en el motor de SQL Server. Las bases de datos de "
     "Fabric son de disponibilidad general desde Ignite 2025."),
    ("Warehouse", "Almacén de datos con T-SQL completo sobre tablas en formato Delta Lake."),
]


@builder
def documents():
    """Documents/: a PDF on the Copilot logos, Spanish meeting notes in Word, and a Spanish Fabric glossary in RTF."""
    body = LOGO_FACTS.replace("{copilot}", asset("microsoft-copilot-2023.svg")).replace(
        "{lockup}", asset("github-copilot-lockup.svg"))
    document = page(body, LOGO_FACTS_CSS).replace("<head>", "<head><title>Copilot logos – quick facts</title>", 1)
    print_pdf(document, out("Documents/Copilot logos – quick facts.pdf"))

    word = f"""<!doctype html><html><head><meta charset="utf-8"><style>body{{font-family:'Helvetica Neue';
font-size:11pt}} h1{{font-size:20pt}} h2{{font-size:13pt;color:#0f4c81}} th{{background:#e8eef6;text-align:left}}
</style></head><body>{MEETING_NOTES}</body></html>"""
    textutil(word, out("Documents/Notas de reunión – piloto de agentes (Contoso).docx"),
             "Notas de reunión – Piloto de agentes (Contoso)")

    entries = "".join(f"<p><b>{esc(term)}</b><br>{esc(text)}</p>" for term, text in GLOSSARY)
    rich = f"""<!doctype html><html><head><meta charset="utf-8"><style>body{{font-family:'Helvetica Neue';
font-size:12pt}} h1{{font-size:22pt}} p{{margin:0 0 10pt}}</style></head><body><h1>Glosario de Microsoft Fabric</h1>
<p><i>Equipo de datos de Contoso · octubre de 2026</i></p>
<p>Microsoft Fabric es una plataforma de análisis de datos de principio a fin, como servicio: ingesta, transformación,
procesamiento en tiempo real, análisis e informes, todo sobre un único lago de datos, OneLake. Se presentó en mayo de
2023 y está disponible con carácter general desde noviembre de 2023.</p>{entries}
<p><i>Fuentes: «What is Microsoft Fabric?», «OneLake shortcuts» y «Mirroring in Microsoft Fabric», en Microsoft Learn
(learn.microsoft.com/fabric). Contoso es una empresa ficticia.</i></p></body></html>"""
    textutil(rich, out("Documents/Glosario de Microsoft Fabric.rtf"), "Glosario de Microsoft Fabric")


# MARK: Slides

FONT, MONO = "Avenir Next", "Menlo"
LEFT, RIGHT = 0.75, 12.583  # the slides' margins, in inches (16:9 is 13.333 × 7.5)


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip("#").upper())


def strong(text: str) -> list[tuple[str, bool]]:
    """"Use **this** one" → its parts, each with whether it's bold."""
    return [(part, index % 2 == 1) for index, part in enumerate(text.split("**")) if part]


def add_run(paragraph, text: str, size: float, color: str, font: str = FONT, bold: bool = False,
            italic: bool = False, tracking: int = 0) -> None:
    run = paragraph.add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold, run.font.italic = font, Pt(size), bold, italic
    run.font.color.rgb = rgb(color)
    if tracking:
        run._r.get_or_add_rPr().set("spc", str(tracking))


def write(frame, items, *, size: float = 16, color: str = "000000", font: str = FONT, bold: bool = False,
          italic: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line: float = 1.12, after: float = 0,
          before: float = 0, tracking: int = 0, bullet: tuple | None = None, pad=(0, 0, 0, 0)):
    """Paragraphs in a text frame: each a string, with "**bold**" parts, or (string, its own style). A bullet is
    (character, color, indent in inches)."""
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left, frame.margin_top, frame.margin_right, frame.margin_bottom = map(Inches, pad)
    defaults = dict(size=size, color=color, font=font, bold=bold, italic=italic, align=align, line=line,
                    after=after, before=before, tracking=tracking, bullet=bullet)
    for index, item in enumerate([items] if isinstance(items, (str, tuple)) else items):
        words, own = (item, {}) if isinstance(item, str) else item
        style = {**defaults, **own}
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = style["align"]
        paragraph.line_spacing = style["line"]
        paragraph.space_before, paragraph.space_after = Pt(style["before"]), Pt(style["after"])
        for part, emphasis in strong(words) or [(" ", False)]:
            add_run(paragraph, part, style["size"], style["color"], style["font"], style["bold"] or emphasis,
                    style["italic"], style["tracking"])
        if style["bullet"]:
            character, tint, indent = style["bullet"]
            properties = paragraph._p.get_or_add_pPr()
            properties.set("marL", str(Inches(indent)))
            properties.set("indent", str(-Inches(indent)))
            etree.SubElement(etree.SubElement(properties, qn("a:buClr")), qn("a:srgbClr")).set("val", tint)
            etree.SubElement(properties, qn("a:buFont")).set("typeface", style["font"])
            etree.SubElement(properties, qn("a:buChar")).set("char", character)
    return frame


def place(shape, x: float, y: float, w: float, h: float) -> None:
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def text(slide, x: float, y: float, w: float, h: float, items, **style):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write(shape.text_frame, items, **style)
    return shape


def box(slide, x: float, y: float, w: float, h: float, fill: str | None = None, stroke: str | None = None,
        weight: float = 1, items=None, **style):
    """A flat rectangle (the theme would give it a shadow), and its text if any. Only square rectangles without
    outlines: Quick Look, so Finder's and Biblo's previews, draws any other shape as a picture and, with several on
    a deck, puts some of those pictures on the wrong slides. An outline of `weight` points is a rectangle in its color
    under a smaller one."""
    if stroke:
        box(slide, x, y, w, h, stroke)
        inset = weight / 72
        x, y, w, h, fill = x + inset, y + inset, w - 2 * inset, h - 2 * inset, fill or "FFFFFF"
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    shape.line.fill.background()
    if items is not None:
        write(shape.text_frame, items, **style)
    return shape


def dot(slide, cx: float, cy: float, side: float, fill: str | None, stroke: str | None = None, **options):
    """A square marker centered on (cx, cy); square for the reason in box()."""
    return box(slide, cx - side / 2, cy - side / 2, side, side, fill, stroke, **options)


def link(slide, x1: float, x2: float, y: float, color: str, weight: float = 1.5, dashed: bool = False) -> None:
    """A connector from x1 to x2 at height y, solid or dashed, ending in a square; rectangles, as in box()."""
    thick = weight / 72
    head = thick * 3.5
    end = x2 - head
    dash, gap = (thick * 4, thick * 2.5) if dashed else (end - x1, 0)
    x = x1
    while x < end - 1e-6:
        box(slide, x, y - thick / 2, min(dash, end - x), thick, color)
        x += dash + gap
    box(slide, end, y - head / 2, head, head, color)


def picture(slide, path: Path, x: float, y: float, w: float | None = None, h: float | None = None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None,
                                    Inches(h) if h else None)


def mark(svg: str) -> tuple[Path, float]:
    """An official logo as a PNG cut to its edges (scaled evenly, nothing else), and its height over its width."""
    image = Image.open(logo(svg, 1024, 0.9))
    image = image.crop(image.getbbox())
    path = WORK / f"mark-{Path(svg).stem}.png"
    image.save(path, optimize=True)
    return path, image.height / image.width


def code(slide, x: float, y: float, w: float, h: float, source: str, size: float = 12) -> None:
    """Python as an editor colors it (Pygments' One Dark), in a dark box."""
    theme = get_style_by_name("one-dark")
    shape = box(slide, x, y, w, h, fill=theme.background_color)
    lines: list[list[tuple[str, str, bool]]] = [[]]
    for kind, value in lex(source, PythonLexer()):
        look = theme.style_for_token(kind)
        for index, piece in enumerate(value.split("\n")):
            if index:
                lines.append([])
            if piece:
                lines[-1].append((piece, look["color"] or "ABB2BF", bool(look["bold"])))
    while lines and not lines[-1]:
        lines.pop()
    frame = shape.text_frame
    write(frame, [], pad=(0.34, 0.3, 0.3, 0.3))
    for index, runs in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 1.16
        for piece, color, bold in runs or [("\u00a0", "ABB2BF", False)]:  # Quick Look drops lines of plain spaces
            add_run(paragraph, piece, size, color, MONO, bold)


class Deck:
    """A 16:9 deck where every slide has a real title (PowerPoint's outline, accessibility) and speaker notes."""

    def __init__(self, palette: dict[str, str], footer: str):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(12192000), Emu(6858000)
        self.p = palette
        self.footer = footer
        self.titles: list[str] = []

    def slide(self, title: str, kicker: str | None = None, notes: str = "", dark: bool = False):
        """A slide. With a kicker it gets the usual heading and footer; without one (covers) its title is placed
        by whoever made it."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        self.titles.append(title)
        if dark:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = rgb(self.p["ink"])
        if kicker is not None:
            p = self.p
            text(slide, LEFT, 0.55, RIGHT - LEFT, 0.3, kicker.upper(), size=11, bold=True, color=p["accent"],
                 tracking=150)
            place(slide.shapes.title, LEFT, 0.84, RIGHT - LEFT, 0.8)
            write(slide.shapes.title.text_frame, title, size=32, bold=True, color=p["ink"], line=1.0)
            text(slide, LEFT, 6.98, 9, 0.25, self.footer, size=9.5, color=p["muted"])
            text(slide, RIGHT - 1, 6.98, 1, 0.25, str(len(self.titles)), size=9.5, color=p["muted"],
                 align=PP_ALIGN.RIGHT)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def save(self, path: Path, title: str, created: dt.datetime, modified: dt.datetime, minutes: int,
             revision: int) -> None:
        """Saves it as PowerPoint would: its own properties (not python-pptx's template's), a thumbnail of its
        first slide, its theme's fonts."""
        properties = self.prs.core_properties
        properties.title, properties.subject, properties.keywords = title, "", ""
        properties.comments, properties.category = "", ""
        properties.author = properties.last_modified_by = "Contoso"
        properties.revision = revision
        properties.created, properties.modified = created, modified  # UTC
        root = self.prs._element
        root.find(qn("p:sldSz")).attrib.pop("type", None)  # the template's "screen4x3", on a 16:9 deck
        if root.find(qn("p:notesMasterIdLst")) is None and any(s.has_notes_slide for s in self.prs.slides):
            # python-pptx relates the notes master without listing it, and macOS (Quick Look, so Biblo's previews
            # too) can't open a deck whose notes master isn't listed.
            master = self.prs.part.relate_to(self.prs.part.notes_master_part, RT.NOTES_MASTER)
            listing = etree.Element(qn("p:notesMasterIdLst"))
            etree.SubElement(listing, qn("p:notesMasterId")).set(qn("r:id"), master)
            root.find(qn("p:sldMasterIdLst")).addnext(listing)
        self.prs.save(path)

        words = paragraphs = 0
        fonts = {FONT}
        for slide in self.prs.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    fonts.update(run.font.name for run in paragraph.runs if run.font.name)
                    said = "".join(run.text for run in paragraph.runs).strip()
                    if said:
                        paragraphs += 1
                        words += len(said.split())
        parts = [*sorted(fonts), "Office Theme", *self.titles]
        listed = "".join(f"<vt:lpstr>{esc(part)}</vt:lpstr>" for part in parts)
        application = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" \
xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><TotalTime>{minutes}</TotalTime>\
<Words>{words}</Words><Application>Microsoft Macintosh PowerPoint</Application>\
<PresentationFormat>Widescreen</PresentationFormat><Paragraphs>{paragraphs}</Paragraphs>\
<Slides>{len(self.titles)}</Slides><Notes>{len(self.titles)}</Notes><HiddenSlides>0</HiddenSlides><MMClips>0</MMClips>\
<ScaleCrop>false</ScaleCrop><HeadingPairs><vt:vector size="6" baseType="variant"><vt:variant><vt:lpstr>Fonts Used\
</vt:lpstr></vt:variant><vt:variant><vt:i4>{len(fonts)}</vt:i4></vt:variant><vt:variant><vt:lpstr>Theme</vt:lpstr>\
</vt:variant><vt:variant><vt:i4>1</vt:i4></vt:variant><vt:variant><vt:lpstr>Slide Titles</vt:lpstr></vt:variant>\
<vt:variant><vt:i4>{len(self.titles)}</vt:i4></vt:variant></vt:vector></HeadingPairs><TitlesOfParts>\
<vt:vector size="{len(parts)}" baseType="lpstr">{listed}</vt:vector></TitlesOfParts><LinksUpToDate>false\
</LinksUpToDate><SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000\
</AppVersion></Properties>""".encode()
        theme = re.compile(rb'(<a:(?:major|minor)Font><a:latin typeface=")[^"]*')
        repack(path, {"docProps/app.xml": application,
                      "ppt/theme/theme1.xml": lambda data: theme.sub(rb"\g<1>" + FONT.encode(), data)})

        # The template's thumbnail is a blank 4:3 page; Quick Look draws the real first slide.
        try:
            subprocess.run(["qlmanage", "-t", "-s", "512", "-o", str(WORK), str(path)], capture_output=True,
                           timeout=60)
        except subprocess.TimeoutExpired:
            pass
        drawn = WORK / f"{path.name}.png"
        if drawn.exists():
            image = Image.open(drawn).convert("RGB")
            image.thumbnail((256, 256))
            buffer = io.BytesIO()
            image.save(buffer, "JPEG", quality=85)
            repack(path, {"docProps/thumbnail.jpeg": buffer.getvalue()})
            drawn.unlink()
        else:
            print(f"  (no thumbnail for {path.name}: Quick Look didn't draw it)", flush=True)


def repack(path: Path, changes: dict) -> None:
    """Writes a zip package again with some of its parts changed (to bytes, or by a function of their bytes), in the
    same order, with the fixed dates Office gives them."""
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    temporary = path.with_name(path.name + ".part")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            if (change := changes.get(name)) is not None:
                data = change(data) if callable(change) else change
            target.writestr(zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0)), data, zipfile.ZIP_DEFLATED)
    temporary.replace(path)


FABRIC_NOTES = [
    "Welcome! In the next fifteen minutes: what Microsoft Fabric is, how OneLake keeps a single copy of our data, "
    "and a short live demo with Contoso sample data: orders and returns in one lakehouse, a Power BI report in "
    "Direct Lake mode, and a summary written by Copilot in Fabric.",
    "Microsoft Learn's definition, word for word. Three things to remember: Fabric is delivered as a service, "
    "everything sits on OneLake, and AI is built in. Copilot in Fabric helps author queries, pipelines and code, "
    "and Fabric works with Microsoft Foundry for models and AI tooling.",
    "The nine workloads Microsoft Learn lists today, all on the same data in OneLake. For this team the "
    "interesting ones are Data Factory, Data Engineering, Power BI and Real-Time Intelligence. Fabric IQ is still "
    "in preview.",
    "OneLake is one logical data lake for the whole tenant, built on ADLS Gen2. Shortcuts point to data stored "
    "somewhere else and behave like symbolic links: that's how we'll read the order history without copying it. "
    "Mirroring continuously replicates databases into OneLake; mirroring for Google BigQuery is generally "
    "available.",
    "Fabric went from preview at Build in May 2023 to general availability at Ignite in November 2023. "
    "Real-Time Intelligence arrived in May 2024 and Fabric Databases in November 2024. At Ignite 2025 databases "
    "became generally available and Fabric IQ entered preview. At GA, 25,000 organizations were already using "
    "Fabric, including 67% of the Fortune 500.",
    "The demo: orders mirrored from Azure SQL Database and the order history through a shortcut, both in the "
    "contoso_sales lakehouse; a Power BI report in Direct Lake mode with sales and returns by region; and Copilot "
    "in Fabric summarizing the quarter. Contoso is fictitious and all of the data is sample data.",
    "Links to everything in this deck. The Fabric icon comes from the Microsoft Fabric icons page on Microsoft "
    "Learn and is shown unaltered.",
]


def fabric_deck() -> None:
    p = {"ink": "0B1F1C", "accent": "0E7C6B", "soft": "E8F4F0", "line": "CFE3DC", "muted": "5E716C"}
    deck = Deck(p, "Microsoft Fabric – platform overview · Contoso data team")
    notes = iter(FABRIC_NOTES)

    s = deck.slide("Microsoft Fabric", notes=next(notes))
    box(s, 0, 0, 5.2, 7.5, fill=p["soft"])
    path, ratio = mark("microsoft-fabric.svg")
    w = min(2.6, 3.0 / ratio)
    picture(s, path, (5.2 - w) / 2, (7.5 - w * ratio) / 2, w, w * ratio)
    text(s, 6.0, 2.1, 6.6, 0.3, "CONTOSO DATA TEAM · SALES DEMO", size=12, bold=True, color=p["accent"],
         tracking=200)
    place(s.shapes.title, 6.0, 2.45, 6.6, 0.95)
    write(s.shapes.title.text_frame, "Microsoft Fabric", size=54, bold=True, color=p["ink"], line=1.0)
    text(s, 6.0, 3.48, 6.0, 0.8, "Platform overview: one copy of our data, from ingestion to Power BI.", size=21,
         color=p["muted"], line=1.18)
    box(s, 6.0, 4.55, 0.8, 0.055, fill=p["accent"])
    text(s, 6.0, 4.8, 6.6, 0.7, [("Thursday, October 8, 2026", {"bold": True, "color": p["ink"]}),
                                 "Javier · Contoso data team"], size=15, color=p["muted"], after=3)

    s = deck.slide("What is Microsoft Fabric?", "What it is", next(notes))
    box(s, LEFT, 1.95, RIGHT - LEFT, 1.8, fill=p["soft"])
    box(s, LEFT, 1.95, 0.08, 1.8, fill=p["accent"])
    text(s, LEFT + 0.45, 2.17, 10.9, 1.45, [
        ("“An analytics platform that supports end-to-end data workflows, including data ingestion, "
         "transformation, real-time stream processing, analytics, and reporting.”", {}),
        ("— “What is Microsoft Fabric?”, Microsoft Learn", {"size": 12, "color": p["muted"], "before": 9})],
         size=21, color=p["ink"], line=1.2)
    cards = [("Software as a service", "Fabric runs as a service: teams sign in and start working, with no clusters "
                                      "or servers to set up first."),
             ("OneLake underneath", "One data lake for the whole organization, built on ADLS Gen2. Every workload "
                                   "reads and writes the same data."),
             ("AI built in", "Copilot in Fabric helps write queries, pipelines and code, and summarizes insights. "
                             "Fabric also works with Microsoft Foundry.")]
    width = (RIGHT - LEFT - 0.6) / 3
    for index, (title, body) in enumerate(cards):
        box(s, LEFT + index * (width + 0.3), 4.1, width, 2.5, "FFFFFF", p["line"],
            items=[(f"0{index + 1}", {"size": 12, "color": p["accent"], "bold": True, "tracking": 100, "after": 7}),
                   (title, {"size": 19, "color": p["ink"], "bold": True, "after": 7}), body],
            size=14, color=p["muted"], line=1.2, pad=(0.3, 0.28, 0.3, 0.2))

    s = deck.slide("Nine workloads, one platform", "What's inside", next(notes))
    workloads = [("Power BI", "Reports and dashboards for everyone in the organization."),
                 ("Data Factory", "Ingest and transform data with pipelines and dataflows."),
                 ("Data Engineering", "Apache Spark and notebooks over the lakehouse."),
                 ("Data Warehouse", "Full T-SQL on tables in the open Delta Lake format."),
                 ("Data Science", "Build, train and track machine learning models."),
                 ("Real-Time Intelligence", "Stream, analyze and act on events as they happen."),
                 ("Databases", "SQL database and Cosmos DB, built into Fabric."),
                 ("Industry Solutions", "Ready-made data solutions for specific industries."),
                 ("Fabric IQ", "A shared semantic layer for business operations. In preview.")]
    height = 1.45
    for index, (name, line) in enumerate(workloads):
        row, column = divmod(index, 3)
        box(s, LEFT + column * (width + 0.3), 1.95 + row * (height + 0.22), width, height, p["soft"],
            items=[(name, {"size": 17, "color": p["ink"], "bold": True, "after": 5}), line], size=13,
            color=p["muted"], line=1.18, pad=(0.28, 0.25, 0.28, 0.2))

    s = deck.slide("OneLake: one copy of your data", "How it works", next(notes))
    text(s, LEFT, 1.95, 4.9, 4.7, [
        "**One lake for the whole organization.** A single logical data lake per tenant, built on ADLS Gen2.",
        "**No copies between tools.** Every workload reads and writes the same data in OneLake.",
        "**Shortcuts** point to data stored somewhere else. They behave like symbolic links.",
        "**Mirroring** continuously replicates databases into OneLake, Google BigQuery included."],
         size=16, color=p["ink"], line=1.22, after=16, bullet=("•", p["accent"], 0.28))
    box(s, 9.0, 1.95, RIGHT - 9.0, 4.65, p["soft"])
    text(s, 9.3, 2.15, 3.0, 0.65, [("OneLake", {"size": 19, "bold": True, "color": p["ink"]}),
                                   "one copy, every workload"], size=11.5, color=p["muted"], after=1)
    for index, (name, detail) in enumerate([("Lakehouse", "contoso_sales · Delta tables"),
                                            ("Warehouse", "T-SQL on the same tables"),
                                            ("Semantic model", "Power BI in Direct Lake mode")]):
        box(s, 9.3, 3.0 + index * 1.13, RIGHT - 9.6, 0.92, "FFFFFF", p["line"],
            items=[(name, {"size": 15, "bold": True, "color": p["ink"], "after": 2}), detail], size=11.5,
            color=p["muted"], anchor=MSO_ANCHOR.MIDDLE, pad=(0.22, 0.08, 0.15, 0.08))
    for top, name, detail, how, after, dashed in [(2.8, "Orders", "Azure SQL Database", "Mirroring", "kept in sync",
                                                   False),
                                                  (4.75, "Order history", "ADLS Gen2 storage", "Shortcut",
                                                   "read in place", True)]:
        box(s, 6.05, top, 1.95, 1.05, "FFFFFF", p["line"],
            items=[(name, {"size": 15, "bold": True, "color": p["ink"], "after": 2}), detail], size=11.5,
            color=p["muted"], anchor=MSO_ANCHOR.MIDDLE, pad=(0.2, 0.08, 0.12, 0.08))
        link(s, 8.06, 8.94, top + 0.525, p["accent"], 1.75, dashed)
        text(s, 8.0, top + 0.13, 1.0, 0.3, how, size=11, bold=True, color=p["accent"], align=PP_ALIGN.CENTER)
        text(s, 8.0, top + 0.62, 1.0, 0.3, after, size=10, color=p["muted"], align=PP_ALIGN.CENTER)

    s = deck.slide("From preview to platform", "Timeline", next(notes))
    milestones = [("May 2023", "Build 2023", "Public preview of “an end-to-end, unified analytics platform”."),
                  ("Nov 2023", "Ignite 2023", "Generally available. Copilot in Fabric enters public preview."),
                  ("May 2024", "Build 2024", "Real-Time Intelligence: streaming data and events in one workload."),
                  ("Nov 2024", "Ignite 2024", "Fabric Databases in preview, starting with SQL database."),
                  ("Nov 2025", "Ignite 2025", "Databases generally available, with Cosmos DB. Fabric IQ in "
                                              "preview.")]
    step = (RIGHT - LEFT - 2.1) / 4
    box(s, LEFT + 0.1, 3.04, step * 4, 0.022, fill=p["line"])
    for index, (when, event, what) in enumerate(milestones):
        x = LEFT + index * step
        text(s, x, 1.95, step - 0.2, 0.3, event.upper(), size=11, bold=True, color=p["accent"], tracking=120)
        text(s, x, 2.24, step - 0.2, 0.5, when, size=24, bold=True, color=p["ink"])
        dot(s, x + 0.1, 3.05, 0.2, p["accent"] if index == len(milestones) - 1 else "FFFFFF", p["accent"], weight=2)
        text(s, x, 3.35, min(step - 0.25, RIGHT - x), 1.3, what, size=13.5, color=p["muted"], line=1.22)
    box(s, LEFT, 5.15, RIGHT - LEFT, 1.4, p["soft"])
    text(s, LEFT + 0.45, 5.15, 2.2, 1.4, "67%", size=46, bold=True, color=p["accent"], anchor=MSO_ANCHOR.MIDDLE)
    text(s, LEFT + 2.7, 5.15, 8.6, 1.4, "of the Fortune 500 were already using Microsoft Fabric when it became "
                                         "generally available in November 2023: 25,000 organizations in all.",
         size=17, color=p["ink"], line=1.22, anchor=MSO_ANCHOR.MIDDLE)

    s = deck.slide("Today's demo", "Thursday, October 8", next(notes))
    steps = [("Bring the data in", "Orders mirrored from Azure SQL Database and the order history through a "
                                   "OneLake shortcut, together in the **contoso_sales** lakehouse."),
             ("Report in Direct Lake", "A Power BI report reads the lakehouse tables directly, without importing a "
                                       "copy: sales and returns by region."),
             ("Ask Copilot", "Copilot in Fabric summarizes the quarter's sales and returns and points out what "
                             "changed.")]
    for index, (title, body) in enumerate(steps):
        x = LEFT + index * (width + 0.3)
        box(s, x, 1.95, width, 3.55, "FFFFFF", p["line"])
        dot(s, x + 0.62, 2.62, 0.64, p["accent"], items=str(index + 1), size=20, bold=True, color="FFFFFF",
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.3, 3.2, width - 0.6, 2.2, [(title, {"size": 19, "bold": True, "color": p["ink"], "after": 8}),
                                                  body], size=14.5, color=p["muted"], line=1.22)
    box(s, LEFT, 5.8, RIGHT - LEFT, 0.8, p["soft"],
        items="Contoso is a fictitious company. Everything in this demo is sample data.", size=14,
        color=p["ink"], anchor=MSO_ANCHOR.MIDDLE, pad=(0.35, 0, 0.35, 0))

    s = deck.slide("Sources", "Further reading", next(notes))
    sources = [("What is Microsoft Fabric?", "learn.microsoft.com/fabric/fundamentals/microsoft-fabric-overview"),
               ("OneLake shortcuts", "learn.microsoft.com/fabric/onelake/onelake-shortcuts"),
               ("Mirroring in Microsoft Fabric", "learn.microsoft.com/fabric/mirroring/overview"),
               ("Microsoft Fabric, now generally available", "Microsoft Fabric blog, November 15, 2023"),
               ("Real-Time Intelligence in Microsoft Fabric", "Microsoft Fabric blog, May 21, 2024"),
               ("Ignite 2025 Book of News", "news.microsoft.com/ignite-2025-book-of-news"),
               ("Microsoft Fabric icons", "learn.microsoft.com/fabric/fundamentals/icons · shown unaltered")]
    for index, (name, where) in enumerate(sources):
        top = 1.95 + index * 0.64
        text(s, LEFT, top + 0.16, 4.9, 0.35, name, size=15, bold=True, color=p["ink"])
        text(s, LEFT + 5.1, top + 0.18, RIGHT - LEFT - 5.1, 0.35, where, size=14, color=p["muted"])
        box(s, LEFT, top + 0.63, RIGHT - LEFT, 0.012, fill=p["line"])

    deck.save(out("Slides/Microsoft Fabric – platform overview.pptx"), "Microsoft Fabric – platform overview",
              dt.datetime(2026, 10, 5, 8, 12), dt.datetime(2026, 10, 7, 16, 40), 96, 14)


AGENT_CODE = '''from azure.ai.projects import AIProjectClient
from azure.ai.projects.models import PromptAgentDefinition
from azure.identity import DefaultAzureCredential

credential = DefaultAzureCredential()
project = AIProjectClient(endpoint=ENDPOINT, credential=credential)
agent = project.agents.create_version(
    agent_name="returns-assistant",
    definition=PromptAgentDefinition(
        model=MODEL, instructions=INSTRUCTIONS, tools=TOOLS),
)

openai = project.get_openai_client()
conversation = openai.conversations.create()
response = openai.responses.create(
    conversation=conversation.id,
    input="¿Cuántos días tengo para devolver algo?",
    extra_body={"agent_reference": {"name": agent.name,
                                    "type": "agent_reference"}},
)
print(response.output_text)
'''

WORKSHOP_NOTES = [
    "Bienvenida. El objetivo del taller es que cada persona del equipo salga con un agente creado, probado y "
    "evaluado. Traed el portátil con todo lo de la última diapositiva ya instalado.",
    "Tres laboratorios de unos 45 minutos. El café es a las 11:15. Las preguntas que no dé tiempo a responder las "
    "recogemos en el canal del piloto.",
    "Un agente no es solo un chat: combina instrucciones, un modelo y herramientas para completar tareas. El "
    "ejemplo es real, del piloto: la consulta de la tienda Trailhead 2 que llegó con una varilla rota.",
    "Microsoft Foundry reúne agentes, modelos y herramientas en un mismo recurso de Azure, con trazas, supervisión "
    "y evaluaciones integradas. Se llamaba Azure AI Studio en 2023, Azure AI Foundry desde noviembre de 2024 y "
    "Microsoft Foundry desde Ignite 2025.",
    "Así está montado el piloto: el chat web llama al agente returns-assistant del proyecto proj-support. El agente "
    "busca en los PDF de la política de devoluciones con búsqueda en archivos, que los guarda en un almacén de "
    "vectores del proyecto, consulta los pedidos en el lakehouse de Fabric con lookup_order y abre la devolución "
    "con create_return. Cada ejecución queda trazada en Application Insights.",
    "Lab 1, en el portal. El modelo lo elegimos juntos. No hace falta tocar la red ni la identidad: proj-support "
    "ya está preparado.",
    "Lab 2: el mismo agente desde Python, con azure-ai-projects 2.x. Nos autenticamos con Entra ID "
    "(DefaultAzureCredential), así que no hay claves en el código. create_version guarda una versión nueva del "
    "agente cada vez, y responses.create conversa con él mediante la API de Responses.",
    "Lab 3: evaluamos con más de cien conversaciones en inglés y en español. Todas las métricas pasan salvo "
    "groundedness, 3,8 sobre un objetivo de 4, lo mismo que vimos en la reunión del 6 de octubre. El ejercicio es "
    "encontrar el porqué en las trazas y corregirlo.",
    "Revisad esta lista antes del miércoles, por favor. El acceso a proj-support lo gestiona Marc. Contoso es una "
    "empresa ficticia y todos los datos del taller son de ejemplo.",
]


def workshop_deck() -> None:
    p = {"ink": "12122A", "accent": "4F46E5", "soft": "EEF0FE", "line": "D9DCF7", "muted": "5D6078"}
    deck = Deck(p, "Taller práctico: agentes con Microsoft Foundry · Contoso")
    notes = iter(WORKSHOP_NOTES)
    bullet = ("•", p["accent"], 0.28)

    s = deck.slide("Taller práctico: agentes con Microsoft Foundry", notes=next(notes), dark=True)
    box(s, 1.0, 2.2, 3.1, 3.1, "FFFFFF")
    path, ratio = mark("microsoft-foundry.svg")
    w = min(1.9, 1.9 / ratio)
    picture(s, path, 1.0 + (3.1 - w) / 2, 2.2 + (3.1 - w * ratio) / 2, w, w * ratio)
    text(s, 4.75, 2.0, 7.8, 0.3, "CONTOSO · EQUIPO DE SOPORTE", size=12, bold=True, color="A5B4FC", tracking=200)
    place(s.shapes.title, 4.75, 2.38, 7.8, 1.3)
    write(s.shapes.title.text_frame, "Taller práctico: agentes con Microsoft Foundry", size=42, bold=True,
          color="FFFFFF", line=1.02)
    text(s, 4.75, 3.88, 7.6, 0.5, "Crea, prueba y evalúa tu primer agente en una mañana.", size=20, color="C7CBE8")
    box(s, 4.75, 4.55, 0.8, 0.055, fill="818CF8")
    text(s, 4.75, 4.8, 7.8, 0.7, [("Miércoles, 14 de octubre de 2026 · 10:00–13:00 · Sala Ártico",
                                   {"bold": True, "color": "FFFFFF"}), "Organiza: Lucía (producto)"],
         size=15, color="A3A7C9", after=3)

    s = deck.slide("Agenda", "Miércoles, 14 de octubre", next(notes))
    agenda = [("10:00", "Introducción", "Qué es un agente y qué ofrece Microsoft Foundry"),
              ("10:30", "Lab 1", "Tu primer agente, en el portal"),
              ("11:15", "Pausa", "Café"),
              ("11:30", "Lab 2", "El mismo agente, desde código (Python)"),
              ("12:15", "Lab 3", "Evaluar antes de publicar"),
              ("12:45", "Preguntas", "Y próximos pasos del piloto")]
    for index, (when, what, detail) in enumerate(agenda):
        top, quiet = 1.95 + index * 0.78, what == "Pausa"
        text(s, LEFT, top + 0.17, 1.4, 0.45, when, size=20, bold=True, color=p["muted"] if quiet else p["accent"])
        text(s, LEFT + 1.6, top + 0.17, 3.0, 0.45, what, size=20, bold=True,
             color=p["muted"] if quiet else p["ink"])
        text(s, LEFT + 4.7, top + 0.22, RIGHT - LEFT - 4.7, 0.45, detail, size=16, color=p["muted"])
        box(s, LEFT, top + 0.77, RIGHT - LEFT, 0.012, fill=p["line"])

    s = deck.slide("¿Qué es un agente?", "Introducción", next(notes))
    text(s, LEFT, 1.9, RIGHT - LEFT, 0.8, "Un agente combina un **modelo** con **instrucciones** y **herramientas** "
                                          "para completar tareas, no solo para responder preguntas.",
         size=19, color=p["ink"], line=1.2)
    parts = [("Instrucciones", "Qué debe hacer, con qué tono y qué no debe hacer nunca: por ejemplo, prometer "
                               "reembolsos."),
             ("Modelo", "Entiende la petición, razona y decide el siguiente paso: responder o usar una "
                        "herramienta."),
             ("Herramientas", "Le permiten consultar y actuar: búsqueda en archivos, funciones, servidores MCP u "
                              "otros agentes.")]
    width = (RIGHT - LEFT - 0.6) / 3
    for index, (name, body) in enumerate(parts):
        box(s, LEFT + index * (width + 0.3), 2.85, width, 2.15, "FFFFFF", p["line"],
            items=[(name, {"size": 19, "bold": True, "color": p["ink"], "after": 7}), body], size=14.5,
            color=p["muted"], line=1.2, pad=(0.3, 0.28, 0.3, 0.2))
    box(s, LEFT, 5.25, RIGHT - LEFT, 1.4, p["soft"],
        items=[("EJEMPLO DEL PILOTO", {"size": 11, "bold": True, "color": p["accent"], "tracking": 150,
                                       "after": 6}),
               "«Mi tienda Trailhead 2 llegó con una varilla rota»: **lookup_order** encuentra el pedido #48213, la "
               "**búsqueda en archivos** aplica la política de devoluciones y **create_return** abre la devolución "
               "RMA-20931."], size=15, color=p["ink"], line=1.22, anchor=MSO_ANCHOR.MIDDLE, pad=(0.4, 0.1, 0.4, 0.1))

    s = deck.slide("Microsoft Foundry, en una diapositiva", "La plataforma", next(notes))
    text(s, LEFT, 1.85, RIGHT - LEFT, 0.7, "La plataforma de Azure para crear aplicaciones y agentes de IA: agentes, "
                                           "modelos y herramientas en un mismo sitio, con trazas, supervisión y "
                                           "evaluaciones integradas.", size=17, color=p["muted"], line=1.2)
    pieces = [("Modelos", "Más de 10.000 modelos de Microsoft, OpenAI, Anthropic, Meta y otros. El enrutador de "
                          "modelos elige el mejor para cada petición."),
              ("Agent Service", "Agentes de instrucciones, de voz y alojados, con tu propio código. Disponible con "
                                "carácter general desde mayo de 2025."),
              ("Herramientas", "Búsqueda en archivos, funciones, catálogo de herramientas MCP y conexión agente a "
                               "agente (A2A)."),
              ("Foundry IQ", "Una capa de conocimiento gestionada para que los agentes respondan con tus datos."),
              ("Guardrails", "Controles sobre las peticiones, las respuestas, las llamadas a herramientas y sus "
                             "resultados."),
              ("Observabilidad", "Trazas, evaluaciones y paneles de supervisión, integrados en cada proyecto.")]
    for index, (name, body) in enumerate(pieces):
        row, column = divmod(index, 3)
        box(s, LEFT + column * (width + 0.3), 2.62 + row * 1.9, width, 1.75, p["soft"],
            items=[(name, {"size": 17, "bold": True, "color": p["ink"], "after": 5}), body], size=13,
            color=p["muted"], line=1.2, pad=(0.28, 0.25, 0.28, 0.2))
    text(s, LEFT, 6.47, RIGHT - LEFT, 0.35, "Se usa desde el portal, los SDK (Python, C#, JavaScript y Java), azd o "
                                           "VS Code con Foundry Toolkit.", size=13, color=p["muted"])

    s = deck.slide("La arquitectura del piloto", "Cómo encaja todo", next(notes))
    architecture = OUT / "Images/RAG architecture – Foundry + Fabric.png"
    if not architecture.exists():
        diagram()
    hairline = 0.75 / 72
    box(s, LEFT - hairline, 1.85 - hairline, 8.0 + 2 * hairline, 4.5 + 2 * hairline, p["line"])
    picture(s, architecture, LEFT, 1.85, 8.0, 4.5)
    text(s, 9.1, 1.95, RIGHT - 9.1, 4.5, [
        "El chat web llama al agente **returns-assistant**, del proyecto **proj-support** (Sweden Central).",
        "Responde con los PDF de la política de devoluciones, mediante búsqueda en archivos, y con los pedidos del "
        "lakehouse de Fabric.",
        "Las guardrails revisan peticiones, respuestas y llamadas a herramientas. Cada ejecución queda trazada."],
         size=15, color=p["ink"], line=1.22, after=16, bullet=bullet)

    s = deck.slide("Lab 1 · Tu primer agente, en el portal", "Lab 1 · 10:30", next(notes))
    steps = ["Abre el proyecto **proj-support** en el portal de Microsoft Foundry.",
             "Crea un agente llamado **returns-assistant** con el modelo que indique el instructor.",
             "Pega sus instrucciones, que están en **prompts/returns-assistant.md**.",
             "Añade **búsqueda en archivos** con los PDF de la política de devoluciones.",
             "Pruébalo en el área de pruebas con la pregunta de la derecha."]
    for index, step in enumerate(steps):
        top = 1.95 + index * 0.9
        dot(s, LEFT + 0.22, top + 0.22, 0.44, p["soft"], items=str(index + 1), size=14, bold=True,
            color=p["accent"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, LEFT + 0.7, top + 0.04, 6.6, 0.8, step, size=16.5, color=p["ink"], line=1.2)
    box(s, 8.4, 1.95, RIGHT - 8.4, 2.3, p["soft"],
        items=[("CONSEJO", {"size": 11, "bold": True, "color": p["accent"], "tracking": 150, "after": 8}),
               "Cada cambio guarda una **versión nueva** del agente. Anota cuál usas: en el Lab 3 compararemos "
               "versiones."], size=15, color=p["ink"], line=1.22, pad=(0.35, 0.3, 0.35, 0.2))
    box(s, 8.4, 4.5, RIGHT - 8.4, 1.75, "FFFFFF", p["line"],
        items=[("PREGUNTA DE PRUEBA", {"size": 11, "bold": True, "color": p["accent"], "tracking": 150,
                                       "after": 8}),
               ("«¿Puedo devolver una tienda de campaña que llegó con una varilla rota?»", {"italic": True})],
        size=15.5, color=p["ink"], line=1.22, pad=(0.35, 0.3, 0.35, 0.2))

    s = deck.slide("Lab 2 · El mismo agente, desde código", "Lab 2 · 11:30", next(notes))
    code(s, LEFT, 1.85, 7.7, 5.0, AGENT_CODE, 11)
    text(s, 8.85, 1.95, RIGHT - 8.85, 4.6, [
        "**AIProjectClient** se conecta al proyecto con tu identidad de Entra ID: sin claves en el código.",
        "**create_version** guarda el agente con su modelo, sus instrucciones y sus herramientas.",
        "**responses.create** conversa con él mediante la API de Responses.",
        "El código completo está en **src/agent.py**."], size=15, color=p["ink"], line=1.22, after=16,
         bullet=bullet)

    s = deck.slide("Lab 3 · Evaluar antes de publicar", "Lab 3 · 12:15", next(notes))
    for x, label in [(LEFT, "MÉTRICA"), (LEFT + 4.0, "OBJETIVO"), (LEFT + 5.5, "ÚLTIMA EJECUCIÓN")]:
        text(s, x, 1.95, 2.0, 0.3, label, size=10.5, bold=True, color=p["muted"], tracking=120)
    box(s, LEFT, 2.32, 7.15, 0.014, fill=p["line"])
    metrics = [("Resolución de la intención", "intent resolution", "≥ 4,0", "4,4", True),
               ("Precisión de las herramientas", "tool call accuracy", "≥ 4,0", "4,2", True),
               ("Cumplimiento de la tarea", "task adherence", "≥ 90 %", "93 %", True),
               ("Fundamentación", "groundedness", "≥ 4,0", "3,8", False),
               ("Relevancia", "relevance", "≥ 4,0", "4,5", True)]
    for index, (name, english, target, result, passed) in enumerate(metrics):
        top = 2.42 + index * 0.8
        text(s, LEFT, top + 0.08, 3.9, 0.65, [name, (english, {"size": 11.5, "color": p["muted"]})], size=16,
             color=p["ink"], line=1.1)
        text(s, LEFT + 4.0, top + 0.12, 1.4, 0.4, target, size=16, color=p["ink"])
        text(s, LEFT + 5.5, top + 0.12, 1.0, 0.4, result, size=16, bold=True,
             color=p["ink"] if passed else "B45309")
        dot(s, LEFT + 6.6, top + 0.3, 0.14, "16A34A" if passed else "F59E0B")
        box(s, LEFT, top + 0.79, 7.15, 0.012, fill=p["line"])
    text(s, 8.45, 1.95, RIGHT - 8.45, 3.0, [
        "Ejecuta **src/evaluate.py** con el conjunto de prueba: más de 100 conversaciones, en inglés y en español.",
        "Abre las trazas de las respuestas con la fundamentación más baja.",
        "Corrige instrucciones o conocimiento, crea una versión nueva y vuelve a evaluar."],
         size=15, color=p["ink"], line=1.22, after=14, bullet=bullet)
    box(s, 8.45, 5.15, RIGHT - 8.45, 1.45, p["soft"],
        items=[("REGLA DEL PILOTO", {"size": 11, "bold": True, "color": p["accent"], "tracking": 150, "after": 6}),
               "No se amplía a ventas hasta llegar a **groundedness 4** o más."], size=15, color=p["ink"],
        line=1.2, anchor=MSO_ANCHOR.MIDDLE, pad=(0.35, 0.1, 0.3, 0.1))

    s = deck.slide("Antes del taller", "Preparación", next(notes))
    checklist = ["VS Code con la extensión **Foundry Toolkit**",
                 "**Python 3.10** o posterior",
                 "**Azure CLI** (az) y **Azure Developer CLI** (azd), con la sesión iniciada",
                 "Acceso al proyecto **proj-support** (pídeselo a Marc)",
                 "El repositorio **contoso-support-agent**, clonado y con sus dependencias instaladas"]
    for index, item in enumerate(checklist):
        top = 2.0 + index * 0.82
        box(s, LEFT, top + 0.05, 0.34, 0.34, None, p["accent"], 1.75)
        text(s, LEFT + 0.6, top, 7.4, 0.5, item, size=18, color=p["ink"])
    box(s, 9.0, 2.0, RIGHT - 9.0, 1.75, p["soft"],
        items=[("¿DUDAS?", {"size": 11, "bold": True, "color": p["accent"], "tracking": 150, "after": 8}),
               "Escribe a Lucía en el canal del piloto. El acceso a **proj-support** lo gestiona Marc."],
        size=15, color=p["ink"], line=1.22, pad=(0.35, 0.3, 0.35, 0.2))
    text(s, LEFT, 6.35, RIGHT - LEFT, 0.35, "Contoso es una empresa ficticia. Todos los datos del taller son de "
                                            "ejemplo.", size=12, italic=True, color=p["muted"])

    deck.save(out("Slides/Taller práctico – agentes con Microsoft Foundry.pptx"),
              "Taller práctico: agentes con Microsoft Foundry", dt.datetime(2026, 10, 7, 9, 30),
              dt.datetime(2026, 10, 9, 11, 15), 143, 21)


@builder
def slides():
    """Slides/: a Microsoft Fabric overview in English, with the Fabric logo, and a Foundry workshop in Spanish."""
    fabric_deck()
    workshop_deck()


if __name__ == "__main__":
    main()
